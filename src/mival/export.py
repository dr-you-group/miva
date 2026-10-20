"""The results bundle a site sends back: one zip, aggregate values only.

Everything in the zip is built here from the describe and evaluate tables and
the stage manifests listed in this invocation's ``runs.json``. Four rules hold
for every file:

1. **Allowlist.** Only paths matching ``ALLOWLIST`` are written; nothing is
   copied from a run directory as is (evaluate's own figures, drawn from
   predictions, never leave).
2. **Small cells.** Every count from 1 to ``min_cell - 1`` is shown as
   ``"masked"``, and so are the further cells hidden to protect it
   (:mod:`mival.suppression`, over one constraint system per study that joins
   Table 1, metrics, sample size, comparisons, exclusions and the STARD
   boxes), so no masked count can be recovered from the counts shown in any
   of the files. A rate, percent or curve point computed from a masked count
   is NaN.
3. **Figures from masked tables.** Every figure is redrawn from the masked
   tables with :mod:`mival.figures`; a per-record figure (the regression
   scatter) has no table form and is never drawn.
4. **Disclosure check.** :func:`check_no_patient_level` reads back every
   written table, JSON and Markdown file. A record identifier or file-path
   column, an absolute path, or a date (other than the bundle's own
   ``generated_at``) fails the export, and no zip is left behind.

``min_cell`` comes from each study's ``stages.export.min_cell``
(:func:`mival.criteria.min_cell`); with several studies the largest is used.
``evaluate.min_events`` below it is refused: the evaluate stage would then
draw curves for slices this bundle has to mask.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import sys
import tempfile
import warnings
import zipfile
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath, PureWindowsPath
from types import SimpleNamespace
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence, Tuple, Union

from mival import criteria
from mival.suppression import Cells, DisclosureError  # noqa: F401 - DisclosureError re-exported
from mival.metrics.curves import COUNT_KEYS
from mival.stages._curve_rows import (
    CALIBRATION,
    CALIBRATION_HORIZON,
    KM_BY_RISK_GROUP,
    REGRESSION_CALIBRATION,
    RISK_DISTRIBUTION,
    SURVIVAL_THRESHOLD_GRID,
    THRESHOLD_GRID,
)
from mival.stages._predictions import DEV_SPLIT, TEST_SPLIT
from mival.stages.describe import COLUMN_OVERALL, EXCLUSION_GROUPS

PathLike = Union[str, Path]


# ---------------------------------------------------------------------------
# What may leave
# ---------------------------------------------------------------------------

_NAME = r"[A-Za-z0-9][A-Za-z0-9._-]*"

#: Every path the bundle may contain, as full-match regular expressions.
ALLOWLIST: Tuple[str, ...] = (
    r"README\.md",
    rf"table1/{_NAME}\.(csv|xlsx)",
    rf"metrics/{_NAME}_(metrics|comparisons|sample_size)\.csv",
    r"metrics/all_metrics\.xlsx",
    rf"curves/{_NAME}_curves\.csv",
    rf"figures/{_NAME}\.png",
    rf"exclusions/{_NAME}\.csv",
    rf"provenance/{_NAME}_{_NAME}_(manifest|summary)\.json",
    r"provenance/export_manifest\.json",
    r"reporting/tripod_ai_checklist\.md",
    r"reporting/methods_draft\.md",
)

#: Column (or JSON key) names that are record-level by definition.
FORBIDDEN_COLUMNS = frozenset({
    "person_id", "image_occurrence_id", "local_path", "measurement_id",
    "visit_occurrence_id", "condition_occurrence_id", "procedure_occurrence_id",
    "prob", "logit", "pred_value",
})
#: Names that carry an individual date: ``*_date``, ``*_datetime``, birth.
_DATE_NAME = re.compile(r"(^|_)(date|datetime|dob)(_|$)|birth", re.IGNORECASE)
_DATE_VALUE = re.compile(r"\d{4}-\d{2}-\d{2}")
_ABSOLUTE = re.compile(r"^\s*(/|\\\\|[A-Za-z]:[\\/])")
#: An absolute path inside free text: a drive path, or ``/a/`` after a space or quote.
_EMBEDDED_ABSOLUTE = re.compile(r"(?<![A-Za-z0-9])[A-Za-z]:[\\/]|(?:^|[\s\"'(=,\[])/[^\s/]+/")
_MD_FORBIDDEN = re.compile(
    r"\b(" + "|".join(sorted(re.escape(n) for n in FORBIDDEN_COLUMNS if "_" in n)) + r")\b"
)

#: The one date-like value the bundle carries: when it was made.
GENERATED_AT = "generated_at"
EXPORT_MANIFEST = "provenance/export_manifest.json"

HOME_TOKEN = "$MIVAL_HOME"
RUNS_TOKEN = "$MIVAL_RUNS"

#: Riley et al. 2021: a validation set should hold at least this many events.
VALIDATION_EVENTS_TARGET = 100

#: Stages whose manifests go into ``provenance/``. Audit-type stages
#: (misclassify) record individual cases and never leave.
MANIFEST_STAGES = ("retrieve", "profile", "preprocess", "describe", "models", "evaluate")

RUNS_FILE = Path("_run") / "runs.json"
SUMMARY_FILES = {"retrieve": "retrieve_summary.json", "profile": "profile_summary.json"}
REQUIRED_STAGES = ("describe", "evaluate")

WITHHELD = "withheld"

#: How every withheld count is shown: the counts from 1 to ``min_cell - 1``
#: and the cells withheld to protect them.
COMPLEMENT_LABEL = "masked"


def masked_label(min_cell: int) -> str:
    """The one label of a withheld count, small or withheld to protect another.

    A single label on purpose: two labels would tell a reader which hidden
    cells are 1 to ``min_cell - 1``.
    """
    return COMPLEMENT_LABEL


# ---------------------------------------------------------------------------
# Small-cell suppression (rule: mival.suppression)
# ---------------------------------------------------------------------------

def _display(kind: str, min_cell: int) -> str:
    return COMPLEMENT_LABEL


def _is_masked(value) -> bool:
    return isinstance(value, str) and value == COMPLEMENT_LABEL


def _group_positions(frame, group_cols: Sequence[str]) -> List[List[int]]:
    if not group_cols:
        return [list(range(len(frame)))]
    groups: Dict[Tuple[str, ...], List[int]] = {}
    keys = zip(*(frame[c].astype(str).tolist() for c in group_cols))
    for position, key in enumerate(keys):
        groups.setdefault(key, []).append(position)
    return list(groups.values())


def suppress(frame, count_cols: Sequence[str], min_cell: int, group_cols: Sequence[str],
             percent_cols: Optional[Mapping[str, str]] = None, strict: bool = True):
    """Mask small counts, with complementary suppression within each group.

    Each group (rows sharing ``group_cols``) of each column of ``count_cols``
    adds up to a total taken as known, so :mod:`mival.suppression` hides
    every cell with ``0 < count < min_cell`` and then further cells until
    no small hidden cell can be recovered (see the module's audit). Every
    hidden cell is shown ``"masked"``. A percent derived from a hidden count
    (the ``<count>_percent`` column, ``percent`` beside a single count
    column, or ``percent_cols``) becomes NaN. Zero is shown; a cell already
    masked stays masked, and its group's total is then unknown.
    ``strict`` refuses a group that cannot be protected (its known total is
    too small); otherwise the best masking is returned.
    """
    import numpy as np

    out = frame.copy()
    positions = _group_positions(out, list(group_cols))
    for column in count_cols:
        raw = out[column].tolist()
        cells = Cells(name=f"{column} by {', '.join(group_cols) or 'table'}")
        index = [None if _is_masked(value) else cells.add(value) for value in raw]
        for group in positions:
            if any(_is_masked(raw[p]) for p in group):
                continue
            parts = [index[p] for p in group if index[p] is not None]
            if parts:
                total = cells.add(sum(cells.values[p] for p in parts), fixed=True)
                cells.constrain(total, parts)
        hidden = cells.solve(min_cell, strict=strict)
        values = out[column].astype(object).to_numpy()
        mask = np.zeros(len(values), dtype=bool)
        for position, cell in enumerate(index):
            if _is_masked(raw[position]):
                mask[position] = True
            elif cell is not None and cell in hidden:
                mask[position] = True
                values[position] = _display(hidden[cell], min_cell)
        out[column] = values
        percent = (percent_cols or {}).get(column)
        if percent is None:
            if f"{column}_percent" in out.columns:
                percent = f"{column}_percent"
            elif len(count_cols) == 1 and "percent" in out.columns:
                percent = "percent"
        if percent is not None:
            out.loc[mask, percent] = np.nan
    return out


def summarize_json(obj: Any, min_cell: int) -> Any:
    """``obj`` with :func:`suppress`'s rule applied to every dict of integers.

    A dict whose values are all integers (booleans excluded) is one group
    with a known total. When its small hidden cells cannot be protected
    (their known total is too small), the integer ``n_*`` totals of the
    same parent dict are masked as well. Any other integer from 1 to
    ``min_cell - 1`` is masked too, except a ``min_*`` setting.
    """
    label = masked_label(min_cell)

    def is_count(value):
        return isinstance(value, int) and not isinstance(value, bool)

    def all_int(value):
        return isinstance(value, dict) and value and all(is_count(v) for v in value.values())

    def mask_dict(body):
        keys = list(body)
        cells = Cells(name="summary")
        index = [cells.add(body[k]) for k in keys]
        total = cells.add(sum(body.values()), fixed=True)
        cells.constrain(total, index)
        hidden = cells.solve(min_cell, strict=False)
        exposed = bool(cells.unprotected or cells.violations(min_cell, hidden))
        return ({k: (_display(hidden[i], min_cell) if i in hidden else body[k])
                 for k, i in zip(keys, index)}, exposed)

    def walk(value):
        if isinstance(value, list):
            return [walk(v) for v in value]
        if not isinstance(value, dict):
            return value
        if all_int(value):
            return mask_dict(value)[0]
        out, exposed = {}, False
        for key, item in value.items():
            if all_int(item):
                out[key], lone = mask_dict(item)
                exposed = exposed or lone
            else:
                out[key] = walk(item)
        for key, item in out.items():
            if not is_count(item):
                continue
            if (exposed and str(key).startswith("n_")) or (
                    0 < item < min_cell and not str(key).startswith("min_")):
                out[key] = label
        return out

    return walk(obj)


# ---------------------------------------------------------------------------
# One constraint system for Table 1, metrics and comparisons
# ---------------------------------------------------------------------------

_VALUE_COLUMNS = ("value", "se", "ci_lo", "ci_hi", "p_value", "p_holm")
_COUNT_STEMS = tuple(COUNT_KEYS)
_THRESHOLD_STEMS = _COUNT_STEMS + (
    "sensitivity", "specificity", "ppv", "npv", "f1", "mcc", "balanced_accuracy",
)
#: Metrics computed on a regression cut's own events (one count family each).
_CUT_STEMS = ("auroc_cut", "auroc_below")
_ARM_COLUMNS = ("site", "model_id", "training_mode", "recipe_id", "perturbation_id", "label_def")
_KINDS = ("n", "e", "ne")


def _slice_columns(frame) -> List[str]:
    from mival.pipeline.runkey import AXES, REPORT_AXES

    return [c for c in tuple(AXES) + tuple(REPORT_AXES) + ("update", "curve") if c in frame.columns]


def metric_slice_columns(frame) -> List[str]:
    """run_key axes, subgroup, outcome and update: one evaluated slice."""
    return [c for c in _slice_columns(frame) if c != "curve"]


def _blank(out, rows, columns=_VALUE_COLUMNS) -> None:
    import numpy as np

    for column in columns:
        if column in out.columns:
            out[column] = out[column].astype(object)
            out.loc[rows, column] = np.nan


def _label(out, rows, column, label) -> None:
    if column in out.columns and rows.any():
        out[column] = out[column].astype(object)
        out.loc[rows, column] = label


def _few_events(frame, min_cell):
    """``(blank, label_events, label_n)`` row masks for the slice-size rule.

    ``blank``: events or non-events (records, without events) below
    ``min_cell``; the row's estimate is not reportable. ``label_events`` and
    ``label_n``: the counts from 1 to ``min_cell - 1`` (``n_events`` too when
    the non-events are). Zero is shown.
    """
    import pandas as pd

    n = pd.to_numeric(frame["n"], errors="coerce")
    events = pd.to_numeric(frame["n_events"], errors="coerce") if "n_events" in frame.columns \
        else pd.Series(float("nan"), index=frame.index)
    nonevents = n - events
    has_events = events.notna()

    def small(count):
        return (count > 0) & (count < min_cell)

    blank = (has_events & ((events < min_cell) | (nonevents < min_cell))) | (n < min_cell)
    label_events = has_events & (small(events) | small(nonevents) | (small(n) & (events > 0)))
    return blank, label_events, small(n)


def _key_value(value) -> str:
    """A key part that reads the same from parquet, csv and a RunKey (None, NaN, 1.0)."""
    if value is None or isinstance(value, bool):
        return "" if value is None else str(value)
    if isinstance(value, float):
        return "" if value != value else "%g" % value
    return str(value)


def _row_keys(frame, columns) -> List[Tuple[str, ...]]:
    if not columns:
        return [()] * len(frame)
    return [tuple(_key_value(v) for v in row) for row in frame[list(columns)].itertuples(index=False)]


def _axis_of(subgroup: str) -> Optional[str]:
    """``"sex=F"`` -> ``"sex"``; the whole slice has no axis."""
    from mival.stages.evaluate import SUBGROUP_ALL

    if subgroup == SUBGROUP_ALL or "=" not in subgroup:
        return None
    return subgroup.split("=", 1)[0]


def _numbers(frame, column):
    import numpy as np
    import pandas as pd

    if column not in frame.columns:
        return np.full(len(frame), np.nan)
    return pd.to_numeric(frame[column], errors="coerce").to_numpy(dtype=float)


class _Counted:
    """One table's rows and the cells their counts are, inside a shared :class:`Cells`."""

    def __init__(self, frame):
        self.frame = frame.reset_index(drop=True)
        self.n: List[Optional[int]] = [None] * len(self.frame)
        self.e: List[Optional[int]] = [None] * len(self.frame)
        self.ne: List[Optional[int]] = [None] * len(self.frame)


def _family(metric: str) -> str:
    stem = metric.split("@", 1)[0]
    return metric if stem in _CUT_STEMS else "core"


def _add_counted(cells: Cells, table: _Counted, prefix: str, slice_cols, base_extra=()):
    """Cells for every row's n, events and non-events, with ``n = events + non-events``.

    Rows of one slice (``slice_cols`` + ``base_extra``) share one ``n`` cell,
    and rows of one count family share their event cells: every metric of a
    slice is shown or hidden together.
    Returns ``(slice key, family)`` per row.
    """
    frame = table.frame
    n = _numbers(frame, "n")
    e = _numbers(frame, "n_events")
    metrics = frame["metric"].astype(str).tolist() if "metric" in frame.columns else [""] * len(frame)
    keys = _row_keys(frame, list(slice_cols) + list(base_extra))
    families = [_family(m) for m in metrics]
    seen: Dict[Tuple, Tuple[float, float]] = {}
    for i, key in enumerate(keys):
        counts = (n[i], e[i])
        previous = seen.setdefault((key, families[i]), counts)
        if not _same_counts(previous, counts):
            families[i] = metrics[i]   # rows of a slice with their own counts
    out = []
    units = set()
    slice_n: Dict[Tuple, float] = {}
    for i, key in enumerate(keys):
        cell = (prefix, key, families[i])
        # One n per slice for every count family (a cut's metrics count the
        # slice's records too); a family with other records keeps its own.
        first_n = slice_n.setdefault(key, n[i])
        n_key = (prefix, key, "n") if _same_counts((first_n,), (n[i],)) else cell + ("n",)
        table.n[i] = cells.add(n[i], key=n_key)
        if e[i] == e[i]:
            table.e[i] = cells.add(e[i], key=cell + ("e",))
            table.ne[i] = cells.add(n[i] - e[i], key=cell + ("ne",))
            if cell not in units:
                units.add(cell)
                cells.constrain(table.n[i], [table.e[i], table.ne[i]])
        out.append((key, families[i]))
    return out


def _same_counts(a, b) -> bool:
    return all((x == y) or (x != x and y != y) for x, y in zip(a, b))


def _constrain_axes(cells: Cells, table: _Counted, slice_cols, family_of):
    """Whole slice = the levels of each subgroup axis; folds add up; updates are equal."""
    frame = table.frame
    subgroup = frame["subgroup"].astype(str).tolist() if "subgroup" in frame.columns \
        else ["all"] * len(frame)
    others = [c for c in slice_cols if c != "subgroup"]
    base = _row_keys(frame, others)
    groups: Dict[Tuple, Dict[str, int]] = {}
    for i in range(len(frame)):
        groups.setdefault((base[i], family_of[i][1]), {}).setdefault(subgroup[i], i)
    for rows in groups.values():
        whole = rows.get("all")
        if whole is None:
            continue
        by_axis: Dict[str, List[int]] = {}
        for value, i in rows.items():
            axis = _axis_of(value)
            if axis is not None:
                by_axis.setdefault(axis, []).append(i)
        for members in by_axis.values():
            for kind in _KINDS:
                cells.constrain(getattr(table, kind)[whole], [getattr(table, kind)[i] for i in members])

    # Folds of one split add up to a split total the bundle shows elsewhere.
    if "fold" in others:
        no_fold = [c for c in others if c != "fold"]
        keys = _row_keys(frame, no_fold)
        folds: Dict[Tuple, Dict[str, int]] = {}
        for i in range(len(frame)):
            folds.setdefault((keys[i], subgroup[i], family_of[i][1]), {}).setdefault(
                _key_value(frame["fold"].iat[i]), i)
        for rows in folds.values():
            if len(rows) < 2:
                continue
            for kind in _KINDS:
                parts = [getattr(table, kind)[i] for i in rows.values()]
                parts = [p for p in parts if p is not None]
                if parts:
                    total = cells.add(sum(cells.values[p] for p in parts), fixed=True)
                    cells.constrain(total, parts)

    # The same records under another update (a recalibration) or in another
    # arm with the same counts: shown or hidden together.
    if "update" in others:
        no_update = [c for c in slice_cols if c != "update"]
        _equal_rows(cells, table, _row_keys(frame, no_update), family_of)
    arm = [c for c in _ARM_COLUMNS if c in frame.columns]
    if arm:
        rest = [c for c in slice_cols if c not in arm]
        keys = _row_keys(frame, rest)
        for kind in _KINDS:
            column = getattr(table, kind)
            buckets: Dict[Tuple, List[int]] = {}
            for i in range(len(frame)):
                if column[i] is not None:
                    value = cells.values[column[i]]
                    buckets.setdefault((keys[i], family_of[i][1], kind, value), []).append(column[i])
            for members in buckets.values():
                for a, b in zip(members, members[1:]):
                    cells.equal(a, b)


def _equal_rows(cells: Cells, table: _Counted, keys, family_of):
    buckets: Dict[Tuple, List[int]] = {}
    for i, key in enumerate(keys):
        buckets.setdefault((key, family_of[i][1]), []).append(i)
    for rows in buckets.values():
        for a, b in zip(rows, rows[1:]):
            for kind in _KINDS:
                cells.equal(getattr(table, kind)[a], getattr(table, kind)[b])


def _table1_cells(cells: Cells, table) -> Dict[str, Any]:
    """Cells of a Table 1 (long form) and the constraints between its columns.

    Count rows are cells. Each (variable, column) also has its denominator:
    the column's ECG count for a categorical variable, else its own ``n``
    cell (non-missing values), whose difference from the ECG count is a
    hidden-able remainder. Columns add up: overall = dev + test, and test =
    its outcome groups (+ records without the label).
    """
    import pandas as pd

    frame = table.reset_index(drop=True)
    stat = frame["stat"].astype(str).tolist()
    variable = frame["variable"].astype(str).tolist()
    level = frame["level"].astype(str).replace({"nan": ""}).tolist()
    column = frame["column"].astype(str).tolist()
    value = pd.to_numeric(frame["value"], errors="coerce").tolist()
    n = pd.to_numeric(frame["n"], errors="coerce").tolist()

    value_cell: List[Optional[int]] = [None] * len(frame)
    for i in range(len(frame)):
        if stat[i] in _TABLE1_COUNT_STATS:
            value_cell[i] = cells.add(value[i], key=("t", variable[i], level[i], column[i]))
    ecgs = {column[i]: value_cell[i] for i in range(len(frame))
            if variable[i] == "ecgs" and stat[i] == "n_ecgs"}
    persons = {column[i]: value_cell[i] for i in range(len(frame))
               if variable[i] == "persons" and stat[i] == "n_persons"}

    # Every person in a column has at least one of its ECGs.
    for col, person_cell in persons.items():
        cells.at_most(ecgs.get(col), [person_cell])

    n_cell: List[Optional[int]] = [None] * len(frame)
    denominators: Dict[Tuple[str, str], int] = {}
    by_var: Dict[Tuple[str, str], List[int]] = {}
    for i in range(len(frame)):
        if variable[i] in ("ecgs", "persons"):
            n_cell[i] = value_cell[i]
            continue
        by_var.setdefault((variable[i], column[i]), []).append(i)
    for (name, col), rows in by_var.items():
        denominator = n[rows[0]]
        total_ecgs = ecgs.get(col)
        if total_ecgs is not None and denominator == cells.values[total_ecgs]:
            cell = total_ecgs
        elif persons.get(col) is not None and denominator == cells.values[persons[col]]:
            cell = persons[col]
        else:
            cell = cells.add(denominator, key=("t", name, "#n", col))
            cells.constrain(total_ecgs, [cell])
        for i in rows:
            n_cell[i] = cell
        if cell is not None:
            denominators[(name, col)] = cell
        counts = [value_cell[i] for i in rows if stat[i] == "count"]
        if counts:
            cells.constrain(cell, counts)

    # Columns that add up, for every count cell and denominator.
    keyed: Dict[Tuple[str, str], Dict[str, int]] = {}
    for i in range(len(frame)):
        if value_cell[i] is not None:
            keyed.setdefault((variable[i], level[i]), {})[column[i]] = value_cell[i]
    for (name, col), rows in by_var.items():
        if n_cell[rows[0]] is not None:
            keyed.setdefault((name, "#n"), {})[col] = n_cell[rows[0]]
    for columns in keyed.values():
        _add_up(cells, columns.get(COLUMN_OVERALL), [columns.get(DEV_SPLIT), columns.get(TEST_SPLIT)])
        groups = [cell for name, cell in columns.items() if name.startswith(f"{TEST_SPLIT}:")]
        _add_up(cells, columns.get(TEST_SPLIT), groups)
    return {"frame": frame, "value_cell": value_cell, "n_cell": n_cell,
            "denominator": denominators}


def _add_up(cells: Cells, total, parts) -> None:
    """``total = sum(parts)`` when every part is present and they do not exceed it."""
    parts = [p for p in parts if p is not None]
    if total is None or len(parts) < 1:
        return
    if sum(cells.values[p] for p in parts) <= cells.values[total] + 1e-9:
        cells.constrain(total, parts)


def _whole_slices(metrics: _Counted, family_of, slice_cols) -> Dict[Tuple, Dict[str, Any]]:
    """Core-family cells of every slice as shipped (no update), the folds of a split together.

    Returns ``{slice key without fold: {"row": a row, "n"/"e"/"ne": set of cells}}``.
    """
    from mival.stages.evaluate import UPDATE_NONE

    frame = metrics.frame
    keys = _row_keys(frame, [c for c in slice_cols if c != "fold"])
    out: Dict[Tuple, Dict[str, Any]] = {}
    for i in range(len(frame)):
        if family_of[i][1] != "core":
            continue
        if "update" in frame.columns and str(frame["update"].iat[i]) != UPDATE_NONE:
            continue
        entry = out.setdefault(keys[i], {"row": i, "n": set(), "e": set(), "ne": set()})
        for kind in _KINDS:
            cell = getattr(metrics, kind)[i]
            if cell is not None:
                entry[kind].add(cell)
    return out


#: What :func:`label_column_of` answers for the primary outcome of a table without label_def.
_PRIMARY_LABEL = object()


def label_column_of(evaluate_spec: Optional[Mapping], primary_outcome: str = "primary"):
    """``(outcome, label_def) -> label column`` as evaluate resolves it, None when unknown.

    An outcome's configured column, else the arm's own label (``label_<def>``,
    a survival arm's event column). ``evaluate_spec`` None: the default study
    whose one outcome is ``primary_outcome``. A metrics table without a
    label_def axis (``label_def`` ``""``) has one label per outcome, so its
    primary outcome stands for the label Table 1 crosses with.
    """
    from mival.stages.evaluate import _Settings

    spec = dict(evaluate_spec) if evaluate_spec is not None else {"outcomes": {primary_outcome: None}}
    settings = _Settings(spec, (), None)
    # Outcome names as the metrics table writes them: YAML may give a key as a
    # number (``30:``), which evaluate's rows carry as the text "30".
    outcomes = {str(name): value for name, value in settings.outcomes.items()}
    primary = str(settings.primary_outcome)

    def column(outcome, label_def: str) -> Optional[str]:
        outcome = str(outcome)
        if outcome not in outcomes:
            return None
        if not label_def:
            return _PRIMARY_LABEL if outcome == primary else None
        configured = outcomes[outcome]
        return settings.label_column_for(label_def) if configured is None else str(configured)

    return column


#: Columns of :func:`scored_counts`: run_key axes, the subgroup slice, the count.
SCORED_COUNT_COLUMNS = ("subgroup", "n_scored")


def scored_mask(key, frame):
    """Rows of one run_key's slice that evaluate's metrics count once they carry a label.

    The same rule as evaluate's slice ``n`` (``_rows_for_slice``,
    ``_regression_rows``, ``_survival_rows``): a classifier counts every
    labelled row, a regression or survival arm the labelled rows with a
    ``pred_value``. So a scored row outside ``n`` has no label.
    """
    import numpy as np

    from mival.stages.evaluate import is_regression, is_survival

    if is_regression(key) or is_survival(key):
        return frame["pred_value"].notna().to_numpy()
    return np.ones(len(frame), dtype=bool)


def scored_count_rows(predictions, evaluate_spec: Optional[Mapping]) -> List[Dict[str, Any]]:
    """Per run_key and subgroup slice (evaluate's own grouping), the scored ECGs."""
    from mival.stages.evaluate import _group_by_run_key, _Settings, _subgroup_slices

    settings = _Settings(dict(evaluate_spec or {}), (), None)
    rows: List[Dict[str, Any]] = []
    for key, group in _group_by_run_key(predictions):
        base = key.to_dict()
        for subgroup, slice_frame in _subgroup_slices(group, settings.subgroups):
            rows.append(dict(base, subgroup=subgroup,
                             n_scored=int(scored_mask(key, slice_frame).sum())))
    return rows


def _evaluate_inputs(run_dir: Path, study_id: str,
                     roots: Optional["RecordedRoots"] = None) -> Dict[str, Any]:
    """The prediction, cohort and report-axis inputs one evaluate run read, from its manifest.

    The manifest lists the inputs by path in the order of their names
    (``attributes``, ``cohort_split``, ``outcomes``, ``predictions``,
    ``runs``, ``train_log``), without the names. Predictions are the
    directory (or the table with every run_key axis), the cohort split the
    table with ``person_id`` and ``split``; a table listed before it is
    ``attributes``, after it ``outcomes``. Each is checked against the
    checksum the run recorded, so the counts are those of the records
    evaluate scored (``runs.json`` and the train log are not read). A path
    recorded under runs that have moved since is read under today's root
    (``roots``, :class:`RecordedRoots`); the checksum still has to match.
    """
    from mival.pipeline.hashing import sha256_path
    from mival.pipeline.runkey import AXES
    from mival.pipeline.tables import read_table

    manifest = _read_json(Path(run_dir) / "manifest.json")
    found: Dict[str, Any] = {"predictions": None, "cohort_split": None,
                             "attributes": [], "outcomes": []}
    tables = []
    gone: List[Path] = []

    def checked(path, ref):
        if sha256_path(path) != ref.get("sha256"):
            raise ValueError(f"study {study_id}: evaluate input {path} changed after the "
                             "evaluate run; its scored ECG counts would not be evaluate's")
        return path

    for ref in manifest.get("input_artifacts") or ():
        path = roots.path(str(ref.get("path"))) if roots is not None else Path(str(ref.get("path")))
        if path.suffix not in (".parquet", ".csv") and not path.is_dir():
            if not path.suffix and not path.exists():
                gone.append(path)   # perhaps the predictions directory
            continue   # runs.json, train_log: not records
        if not path.exists():
            raise ValueError(f"study {study_id}: evaluate input {path} is gone; the scored ECG "
                             "counts the suppression audit needs cannot be read")
        if path.is_dir():
            found["predictions"] = checked(path, ref)
            continue
        columns = set(read_table(path).columns)
        if set(AXES) <= columns and "prob" in columns:
            found["predictions"] = checked(path, ref)
        else:
            tables.append((checked(path, ref), columns))
    cohort = next((i for i, (_, columns) in enumerate(tables)
                   if {"person_id", "split"} <= columns), None)
    if found["predictions"] is None and gone:
        raise ValueError(f"study {study_id}: evaluate input {gone[0]} is gone; the scored ECG "
                         "counts the suppression audit needs cannot be read")
    if found["predictions"] is None or cohort is None:
        raise ValueError(f"study {study_id}: the evaluate run {run_dir} lists no predictions or "
                         "cohort split among its inputs; scored ECG counts cannot be read")
    found["cohort_split"] = tables[cohort][0]
    found["attributes"] = [path for path, _ in tables[:cohort]]
    found["outcomes"] = [path for path, _ in tables[cohort + 1:]]
    return found


def scored_counts(evaluate_dirs: Sequence[Path], evaluate_spec: Optional[Mapping],
                  study_id: str = "", roots: Optional["RecordedRoots"] = None):
    """Each run_key's scored ECGs per subgroup slice, from the inputs its evaluate run read.

    Read at the site from files the runs already have (the predictions,
    cohort split and report-axis tables evaluate joined, as evaluate joins
    them) and never written to the bundle. The suppression audit uses them
    to know where an arm scored every Table 1 ECG (:func:`_link_table1`).
    """
    import pandas as pd

    from mival.pipeline.runkey import AXES
    from mival.pipeline.tables import read_table
    from mival.stages.evaluate import _Settings, _attach_cohort_columns, _attach_table, \
        _load_predictions

    settings = _Settings(dict(evaluate_spec or {}), (), None)
    rows: List[Dict[str, Any]] = []
    for run_dir in evaluate_dirs:
        found = _evaluate_inputs(Path(run_dir), study_id, roots)
        predictions = _load_predictions(found["predictions"], settings.label_def)
        predictions = _attach_cohort_columns(predictions, read_table(found["cohort_split"]))
        for name in ("outcomes", "attributes"):
            for path in found[name]:
                predictions = _attach_table(predictions, read_table(path), name, [])
        rows.extend(scored_count_rows(predictions, evaluate_spec))
    return pd.DataFrame(rows, columns=list(AXES) + list(SCORED_COUNT_COLUMNS))


def _link_table1(cells: Cells, metrics: _Counted, whole, t1, outcome_variable,
                 primary_outcome, label_column=None, scored_counts=None):
    """A metrics slice counts a subset of Table 1's records of the same split and level.

    Table 1 counts every cohort ECG of a split. Inside each ``(split, level)``
    sit counts no file shows: for each label (outcome x label_def) its
    positive, negative and unlabelled ECGs (for a label without events, its
    labelled and unlabelled ECGs), and for each arm the ECGs it scored. Those
    are latent cells (:meth:`Cells.latent`), which the audit protects like
    shown ones, so a reader who bounds them by Table 1 and the metrics cannot
    pin a small number of unlabelled ECGs (say) down. A count Table 1 does
    show takes the latent cell's place: for ``describe.outcome`` (the primary
    outcome) the whole split's label ``1`` and ``0`` counts, and in a split
    Table 1 crosses with the label (test) each level's ``test:<label>=1`` /
    ``=0`` count.

    When ``scored_counts`` (:func:`scored_counts`) shows that an arm
    scored every ECG of a Table 1 count, that is known to the reader:
    scored = Table 1's count. At a uniform site every arm scores every
    Table 1 ECG and a reader may believe so, so the audit must not rely on
    the scored count being free there. The scored count stays a latent cell
    tied to Table 1's (where Table 1 hides a small level, knowing the arm
    scored all of it does not tell how many there are). A gap of 1 or more
    is a count no file shows and stays unknown (taking it as known would
    hand the reader a small count no table gives). The relations, all true
    by definition:

    - Table 1 ``(axis, level)`` of a split (or the split's ECGs) = positive +
      negative + unlabelled, for every label;
    - an axis's levels add up to the whole split, for every latent kind
      (Table 1 lists every level, unknown included);
    - an arm's scored ECGs <= Table 1's count (=, when it scored them all);
      its slice n <= its scored ECGs;
      its scored ECGs <= the label's unlabelled ECGs + n (a scored ECG not in
      n has no label); events <= positives, non-events <= negatives (or n <=
      labelled).
    """
    from mival.stages.evaluate import SUBGROUP_ALL

    frame = metrics.frame
    t1_frame = t1["frame"]
    t1_stat = t1_frame["stat"].astype(str).tolist()
    t1_variable = t1_frame["variable"].astype(str).tolist()
    t1_column = t1_frame["column"].astype(str).tolist()
    t1_level = t1_frame["level"].astype(str).replace({"nan": ""}).tolist()
    levels: Dict[Tuple[str, str], List[str]] = {}
    for i, cell in enumerate(t1["value_cell"]):
        if cell is not None and t1_stat[i] == "count":
            levels.setdefault((t1_variable[i], t1_column[i]), []).append(t1_level[i])
    columns = set(t1_column)
    slice_cols = metric_slice_columns(frame)
    arm_cols = [c for c in slice_cols
                if c not in ("split", "fold", "subgroup", "outcome", "update")]
    arms = _row_keys(frame, arm_cols)
    binary: Dict[Tuple[str, str], bool] = {}
    made: Dict[Tuple, Optional[int]] = {}
    resolve = label_column if label_column is not None else label_column_of(None, primary_outcome)
    known: Optional[Dict[Tuple, float]] = None
    if scored_counts is not None:
        known = {}
        missing = [c for c in arm_cols + ["split", "subgroup", "n_scored"]
                   if c not in scored_counts.columns]
        if missing:
            raise ValueError(f"scored counts lack the columns {missing}")
        scored_keys = _row_keys(scored_counts, arm_cols)
        for arm, split_value, subgroup, count in zip(
                scored_keys, scored_counts["split"].tolist(), scored_counts["subgroup"].tolist(),
                _numbers(scored_counts, "n_scored")):
            # Dev folds are one split: their scored ECGs add up.
            key = (arm, _key_value(split_value), str(subgroup))
            known[key] = known.get(key, 0.0) + float(count)

    def records(split, axis, level):
        if axis is None:
            return cells.get(("t", "ecgs", "", split))
        return cells.get(("t", axis, level, split))

    def partitioned(split, axis):
        """The axis's Table 1 levels when they add up to the split's ECGs."""
        names = levels.get((axis, split), [])
        total = records(split, None, None)
        if not names or total is None:
            return []
        parts = [records(split, axis, name) for name in names]
        if any(p is None for p in parts) or \
                abs(sum(cells.values[p] for p in parts) - cells.values[total]) > 1e-9:
            return []
        return names

    def shown_part(kind, split, axis, level, label):
        """Table 1's own cell for a label part, when Table 1 shows it."""
        # Only the label scored against describe.outcome's own column has
        # Table 1's label counts as its positives and negatives: a survival
        # arm under the same outcome name counts deaths, not that label.
        column = resolve(*label)
        primary = outcome_variable is not None and (column is _PRIMARY_LABEL
                                                    or column == outcome_variable)
        if not primary or kind not in ("pos", "neg"):
            return None
        flag = "1" if kind == "pos" else "0"
        if axis is None:
            return cells.get(("t", outcome_variable, flag, split))
        if f"{split}:{outcome_variable}={flag}" in columns:
            return cells.get(("t", axis, level, f"{split}:{outcome_variable}={flag}"))
        return None

    def kinds(label):
        return ("pos", "neg", "unl") if binary[label] else ("lab", "unl")

    def label_parts(split, axis, level, label):
        """``{kind: cell}`` of a label within one Table 1 count, made once."""
        key = ("label", split, axis, level, label)
        if key in made:
            return made[key]
        total = records(split, axis, level)
        if total is None:
            made[key] = None
            return None
        parts = {}
        for kind in kinds(label):
            parts[kind] = shown_part(kind, split, axis, level, label)
            if parts[kind] is None:
                parts[kind] = cells.latent((kind, split, axis, level, label))
        made[key] = parts
        cells.define(total, list(parts.values()))
        if axis is not None:
            label_parts(split, None, None, label)
            names = partitioned(split, axis)
            if level in names:
                for name in names:
                    label_parts(split, axis, name, label)
                for kind in kinds(label):
                    members = [made[("label", split, axis, name, label)][kind] for name in names]
                    whole_cell = made[("label", split, None, None, label)][kind]
                    if any(m in cells.latent_cells for m in members + [whole_cell]):
                        cells.define(whole_cell, members)
        return parts

    def scored(arm, split, axis, level):
        """The ECGs one arm scored within one Table 1 count, made once."""
        key = ("scored", arm, split, axis, level)
        if key in made:
            return made[key]
        total = records(split, axis, level)
        if total is None:
            made[key] = None
            return None
        # Not audited itself: a small scored count is bounded by its slice's n
        # and Table 1's count, which are.
        cell = made[key] = cells.latent(key, audited=False)
        cells.at_most(total, [cell])
        subgroup = SUBGROUP_ALL if axis is None else f"{axis}={level}"
        count = None if known is None else known.get((arm, split, subgroup))
        if count is not None:
            unscored = cells.values[total] - count
            if unscored < 0:
                raise ValueError(
                    f"evaluate's scored counts exceed Table 1's ECGs for {arm + (split, subgroup)} "
                    f"({count:g} > {cells.values[total]:g}); they come from another run"
                )
            if unscored == 0:
                # Reader-known: the arm scored every ECG of this Table 1
                # count (a uniform site), so scored = Table 1's count. A
                # zero is public, as a shown 0 is; a gap of 1 or more stays
                # unknown, like any count no file shows.
                cells.define(total, [cell])
        if axis is not None:
            scored(arm, split, None, None)
            names = partitioned(split, axis)
            if level in names:
                members = [scored(arm, split, axis, name) for name in names]
                cells.define(made[("scored", arm, split, None, None)], members)
        return cell

    for entry in whole.values():
        i = entry["row"]
        split = str(frame["split"].iat[i]) if "split" in frame.columns else ""
        subgroup = str(frame["subgroup"].iat[i]) if "subgroup" in frame.columns else "all"
        outcome = str(frame["outcome"].iat[i]) if "outcome" in frame.columns else primary_outcome
        label = (outcome, str(frame["label_def"].iat[i]) if "label_def" in frame.columns else "")
        binary.setdefault(label, bool(entry["e"]))
        axis = _axis_of(subgroup)
        level = subgroup.split("=", 1)[1] if axis is not None else None
        n = sorted(entry["n"])
        if known is not None and (arms[i], split, subgroup) not in known:
            raise ValueError(
                f"evaluate's scored counts have no row for the metrics slice "
                f"{arms[i] + (split, subgroup)}; they come from another evaluate run"
            )
        parts = label_parts(split, axis, level, label)
        mine = scored(arms[i], split, axis, level)
        if parts is None or mine is None:
            continue
        cells.at_most(mine, n)
        # A scored ECG outside n has no label: scored <= unlabelled + n.
        bound = cells.latent(("unl+n", arms[i], split, axis, level, label), audited=False)
        cells.define(bound, [parts["unl"]] + n)
        cells.at_most(bound, [mine])
        if not binary[label]:
            cells.at_most(parts["lab"], n)
        elif entry["e"]:
            cells.at_most(parts["pos"], sorted(entry["e"]))
            cells.at_most(parts["neg"], sorted(entry["ne"]))
        else:
            labelled = cells.latent(("lab", split, axis, level, label), audited=False)
            cells.define(labelled, [parts["pos"], parts["neg"]])
            cells.at_most(labelled, n)


def _link_comparisons(cells: Cells, comparisons: _Counted, metrics: Optional[_Counted],
                      metric_slices: Dict[Tuple, int], slice_cols: Sequence[str]):
    """A comparison counts the records both arms scored: a subset of each arm's slice.

    As :meth:`Cells.at_most`, so an equality (hidden in both or neither) when
    the counts agree.
    """
    from mival.pipeline.runkey import RunKey
    from mival.stages.evaluate import UPDATE_NONE

    if metrics is None:
        return
    frame = comparisons.frame
    for i in range(len(frame)):
        for arm in ("arm_a", "arm_b"):
            try:
                key = RunKey.from_string(str(frame[arm].iat[i])).to_dict()
            except (ValueError, KeyError):
                continue
            key.update(subgroup=frame["subgroup"].iat[i], outcome=frame["outcome"].iat[i],
                       update=UPDATE_NONE)
            slice_key = tuple(_key_value(key.get(c)) for c in slice_cols)
            row = metric_slices.get(slice_key)
            if row is None:
                continue
            for kind in _KINDS:
                cells.at_most(getattr(metrics, kind)[row], [getattr(comparisons, kind)[i]])


def _apply_counted(table: _Counted, hidden, min_cell, cells: Cells):
    """The table with hidden counts labelled and the estimates of hidden rows blanked."""
    import numpy as np
    import pandas as pd

    out = table.frame.copy()
    rows = np.zeros(len(out), dtype=bool)
    n_disp, e_disp, ne_disp = [], [], []
    for i in range(len(out)):
        for kind, target in (("n", n_disp), ("e", e_disp), ("ne", ne_disp)):
            cell = getattr(table, kind)[i]
            if cell is None:
                target.append(None)
                continue
            if cell in hidden:
                rows[i] = True
                target.append(_display(hidden[cell], min_cell) if cells.values[cell] > 0
                              else 0)
            else:
                value = cells.values[cell]
                target.append(int(value) if float(value).is_integer() else value)
    n = _numbers(out, "n")
    e = _numbers(out, "n_events")
    with np.errstate(invalid="ignore"):
        small = (n < min_cell) | ((e == e) & ((e < min_cell) | ((n - e) < min_cell)))
    blank = pd.Series(rows | small, index=out.index)
    _blank(out, blank)
    out["n"] = [d if d is not None else v for d, v in zip(n_disp, out["n"].tolist())]
    if "n_events" in out.columns:
        out["n_events"] = [d if d is not None else v for d, v in zip(e_disp, out["n_events"].tolist())]
        out["n_nonevents"] = ne_disp
    return out, blank


def _mask_2x2(out, masked, min_cell):
    """The 2x2 at one operating point masked together, with the rates computed from it."""
    import pandas as pd

    if "metric" not in out.columns or not len(out):
        return masked
    metric = out["metric"].astype(str)
    stem = metric.str.split("@", n=1).str[0]
    suffix = metric.str.split("@", n=1).str[1].fillna("")
    keys = _slice_columns(out)
    is_count = stem.isin(_COUNT_STEMS)
    counts = pd.to_numeric(out["value"], errors="coerce")
    small = is_count & (counts > 0) & (counts < min_cell) & ~masked
    tables = set()
    for position in out.index[small]:
        tables.add(tuple(str(out.at[position, c]) for c in keys) + (suffix[position],))
    if not tables:
        return masked
    row_key = [tuple(str(v) for v in row) for row in out[keys].itertuples(index=False)]
    in_table = pd.Series([k + (s,) in tables for k, s in zip(row_key, suffix)], index=out.index)
    hit = in_table & stem.isin(_THRESHOLD_STEMS)
    _blank(out, hit & ~is_count)
    _blank(out, hit & is_count, [c for c in _VALUE_COLUMNS if c != "value"])
    labels = [masked_label(min_cell) if 0 < c < min_cell else COMPLEMENT_LABEL for c in counts]
    out["value"] = out["value"].astype(object)
    for position in out.index[hit & is_count]:
        out.at[position, "value"] = labels[position] if counts[position] > 0 else 0
    return masked | hit


def mask_tables(min_cell: int, table1=None, metrics=None, comparisons=None,
                outcome_variable: Optional[str] = None, primary_outcome: str = "primary"):
    """Table 1, metrics and comparisons masked as one constraint system.

    Returns ``(table1, metrics, comparisons, slices)``, each masked table or
    None, and ``slices``: for every metrics slice (run_key axes, subgroup,
    outcome, update) whose core counts are hidden, the shown ``(n,
    n_events)`` labels, so its curves can be masked to match. See
    :func:`mask_study`, which also takes the exclusions and STARD boxes.
    """
    out = mask_study(min_cell, table1=table1, metrics=metrics, comparisons=comparisons,
                     outcome_variable=outcome_variable, primary_outcome=primary_outcome)
    return out.table1, out.metrics, out.comparisons, out.slices


def mask_study(min_cell: int, table1=None, metrics=None, comparisons=None, exclusions=None,
               stard=None, outcome_variable: Optional[str] = None,
               primary_outcome: str = "primary", tolerate_exclusions: bool = False,
               evaluate_spec: Optional[Mapping] = None, scored=None):
    """Every count a study exports, masked as one constraint system.

    The relations between tables (each stated where it is added):

    - Table 1 (:func:`_table1_cells`): levels add up to their denominator,
      overall = dev + test, test = its outcome columns, persons <= ECGs.
    - metrics and comparisons (:func:`_add_counted`, :func:`_constrain_axes`):
      n = events + non-events, the whole slice = each axis's levels, folds add
      up, updates and same-count arms are the same records.
    - metrics within Table 1 (:func:`_link_table1`), through latent cells
      for the counts no file shows (each label's positive, negative and
      unlabelled ECGs and each arm's scored ECGs per split and level; with
      ``scored``, an arm that scored all of a count is known to have), and
      comparisons within their arms' slices (:func:`_link_comparisons`):
      subsets, ``<=``.
    - sample_size shows the metrics' own slice cells (it is read from the
      masked metrics).
    - exclusions (:func:`_exclusion_cells`): levels add up to their reason,
      reasons to the stage total, which is fixed.
    - STARD boxes (:func:`_stard_cells`): each box is the next plus the
      stage's exclusions, the last is analysed + dev, analysed and dev are
      subsets of Table 1's test and dev ECGs and supersets of every whole
      metrics slice of that split.

    ``evaluate_spec`` (the study's evaluate stage) says which label column
    each ``(outcome, label_def)`` was scored against (:func:`label_column_of`),
    so only arms scored against ``outcome_variable`` meet Table 1's label
    counts. ``scored`` is :func:`scored_counts`'s table (each run_key's
    scored ECGs per subgroup slice): with it, wherever an arm scored every
    ECG of a Table 1 count, scored = that count is known to the reader, so a
    reader who believes the arms scored every ECG (true at a uniform site)
    is covered; without it (a caller that has only the tables) the scored
    counts are free, which protects only against a reader who does not
    know them. The export always passes it. ``stard`` is ``(stages,
    exclusion lines, side boxes)`` as evaluate counts them. ``tolerate_exclusions``: small exclusion counts that only a stage
    total could protect do not refuse the export (the site allowed it).

    Returns a namespace: ``table1``, ``metrics``, ``comparisons``,
    ``exclusions`` (masked or None), ``slices`` (see :func:`mask_tables`) and
    ``stard`` (``(stages, side)`` with hidden boxes labelled, or None).
    """
    cells = Cells(name="Table 1, metrics, comparisons, exclusions and STARD")
    t1 = _table1_cells(cells, _check_table1_stats(table1)) if table1 is not None else None
    counted = None
    family_of = []
    slice_cols: List[str] = []
    metric_slices: Dict[Tuple, int] = {}
    whole: Dict[Tuple, Dict[str, Any]] = {}
    if metrics is not None and len(metrics):
        counted = _Counted(metrics)
        slice_cols = metric_slice_columns(counted.frame)
        family_of = _add_counted(cells, counted, "m", slice_cols)
        _constrain_axes(cells, counted, slice_cols, family_of)
        whole = _whole_slices(counted, family_of, slice_cols)
        if t1 is not None:
            _link_table1(cells, counted, whole, t1, outcome_variable, primary_outcome,
                         label_column_of(evaluate_spec, primary_outcome), scored)
        for i, (key, family) in enumerate(family_of):
            if family == "core":
                metric_slices.setdefault(key, i)
    comp = None
    if comparisons is not None and len(comparisons):
        comp = _Counted(comparisons)
        comp_slice = [c for c in ("comparison", "arm_a", "arm_b", "subgroup", "outcome",
                                  "category", "metric", "x") if c in comp.frame.columns]
        comp_family = _add_counted(cells, comp, "c", comp_slice)
        comp_family = [(key, "core") for key, _ in comp_family]
        _constrain_axes(cells, comp, comp_slice, comp_family)
        same = _row_keys(comp.frame, [c for c in ("comparison", "arm_a", "arm_b", "subgroup",
                                                    "outcome") if c in comp.frame.columns])
        events = list(zip(_numbers(comp.frame, "n"), _numbers(comp.frame, "n_events")))
        _equal_rows(cells, comp, [k + tuple(_key_value(v) for v in c) for k, c in zip(same, events)],
                    comp_family)
        _link_comparisons(cells, comp, counted, metric_slices, slice_cols)
    flow = None
    if exclusions is not None:
        flow = _exclusion_cells(cells, exclusions, tolerate_exclusions)
    boxes = None
    if stard is not None:
        boxes = _stard_cells(cells, stard, flow["totals"] if flow else {}, counted, whole)

    # Whole slices, the rows a reader most needs, are shown again when no
    # small cell needs them hidden. Subgroup rows are not tried: showing more
    # of them lets a reader who takes two arms to have scored the same ECGs
    # (a set-level relation this system does not hold; it holds their counts
    # only) pin unlabelled counts.
    has_subgroup = counted is not None and "subgroup" in counted.frame.columns
    for entry in whole.values():
        if not has_subgroup or _axis_of(str(counted.frame["subgroup"].iat[entry["row"]])) is None:
            cells.prefer_together([c for kind in _KINDS for c in entry[kind]])
    # strict: an unresolved constraint or an unprotected small cell refuses the export.
    hidden = cells.solve(min_cell)

    table1_out = _apply_table1(t1, hidden, min_cell, cells) if t1 is not None else None
    metrics_out, slices = None, {}
    if counted is not None:
        metrics_out, blank = _apply_counted(counted, hidden, min_cell, cells)
        metrics_out["masked"] = _mask_2x2(metrics_out, blank, min_cell).to_numpy()
        for i, (key, family) in enumerate(family_of):
            if family == "core" and any(getattr(counted, k)[i] in hidden for k in _KINDS
                                        if getattr(counted, k)[i] is not None):
                slices[key] = (metrics_out["n"].iat[i],
                               metrics_out["n_events"].iat[i] if "n_events" in metrics_out else None)
    elif metrics is not None:
        metrics_out = metrics.copy()
        metrics_out["masked"] = False
    comparisons_out = None
    if comp is not None:
        comparisons_out, blank = _apply_counted(comp, hidden, min_cell, cells)
        comparisons_out["masked"] = blank.to_numpy()
    elif comparisons is not None:
        comparisons_out = comparisons.copy()
        comparisons_out["masked"] = False
    exclusions_out = None
    if flow is not None:
        frame = flow["frame"].copy()
        frame["n"] = [
            _display(hidden[c], min_cell) if c is not None and c in hidden else _int_or(v)
            for c, v in zip(flow["index"], frame["n"].tolist())
        ]
        exclusions_out = frame
    stard_out = None
    if boxes is not None:
        def shown(entries, index):
            return [(name, _display(hidden[c], min_cell) if c is not None and c in hidden
                     and cells.values[c] > 0 else count)
                    for (name, count), c in zip(entries, index)]

        stard_out = (shown(stard[0], boxes["stages"]), shown(stard[2], boxes["side"]))
    return SimpleNamespace(table1=table1_out, metrics=metrics_out, comparisons=comparisons_out,
                           exclusions=exclusions_out, slices=slices, stard=stard_out)


def mask_metrics(metrics, min_cell: int):
    """``metrics_long`` with small counts masked (see :func:`mask_tables`).

    A slice's records, events and non-events are cells with ``n = events +
    non-events``; the whole-slice row is the sum of each subgroup axis's
    levels; folds add up; a recalibrated (``update``) slice and another
    arm's slice with the same counts are the same cells. Every metric row of
    a slice shares its cells (a regression cut's metrics have their own
    event cells, and the slice's ``n``), so all its metrics are shown or
    hidden together. A hidden row's estimates are blanked. A row whose
    events or non-events (records, without events) are below ``min_cell``
    is blanked as well, even when its counts are shown (a zero).

    The 2x2 counts at one operating point (``tp`` ... ``fn``, same ``@``
    suffix) are masked together when any is 1 to ``min_cell - 1``, and so are
    the rates computed from them. ``masked`` marks every touched row and
    ``n_nonevents`` shows the non-event count as the table shows it.
    """
    return mask_tables(min_cell, metrics=metrics)[1]


def mask_comparisons(comparisons, min_cell: int):
    """The comparisons table masked on its own (see :func:`mask_tables`)."""
    return mask_tables(min_cell, comparisons=comparisons)[2]


def _x_key(frame) -> List[Tuple[str, ...]]:
    keys = _slice_columns(frame) + ["x"]
    return [tuple(str(v) for v in row) for row in frame[keys].itertuples(index=False)]


def _count_display(value, min_cell: int):
    """A count of a masked point: ``"masked"`` (see :func:`masked_label`); zero stays 0."""
    try:
        number = float(value)
    except (TypeError, ValueError):
        return value
    if number != number or number == 0:
        return value
    return masked_label(min_cell) if number < min_cell else COMPLEMENT_LABEL


def _mask_points(out, curve: str, bad_metric_rows, min_cell: int, count_metrics=()) -> None:
    """Mask every metric at the (slice, x) points of ``bad_metric_rows`` of one curve."""
    import pandas as pd

    if not bad_metric_rows.any():
        return
    keys = _x_key(out)
    bad = {keys[i] for i in range(len(out)) if bad_metric_rows.iat[i]}
    hit = pd.Series([k in bad for k in keys], index=out.index) & (out["curve"] == curve)
    counts = hit & out["metric"].isin(list(count_metrics))
    values = out["value"].tolist()
    _blank(out, hit & ~counts)
    _blank(out, counts, ("ci_lo", "ci_hi"))
    out["value"] = out["value"].astype(object)
    for position in out.index[counts]:
        out.at[position, "value"] = _count_display(values[position], min_cell)
    out.loc[hit, "masked"] = True


def mask_grid_rows(curves, min_cell: int):
    """``threshold_grid`` points masked whole when any of tp/fp/tn/fn is 1 to ``min_cell - 1``.

    Zero is exact and shown, as in :func:`suppress`: a grid end with no
    predicted positive keeps its (NaN) PPV and is not masked. Every metric
    at such a ``(arm, slice, x)`` point — the counts, the rates and net
    benefit — becomes NaN, and ``masked`` is True for its rows.
    """
    import pandas as pd

    out = curves.copy().reset_index(drop=True)
    if "masked" not in out.columns:
        out["masked"] = False
    grid = out["curve"] == THRESHOLD_GRID
    value = pd.to_numeric(out["value"], errors="coerce")
    bad = grid & out["metric"].isin(_COUNT_STEMS) & (value > 0) & (value < min_cell)
    _mask_points(out, THRESHOLD_GRID, bad, min_cell)
    _blank(out, out["masked"] & grid)
    return out


def mask_curves(curves, min_cell: int, hidden_slices: Optional[Mapping] = None):
    """Every ``curves_long`` count masked, with the points computed from it.

    - A slice with events, non-events or (regression) records below
      ``min_cell`` loses all its points; so does a slice whose counts the
      metrics table hides (``hidden_slices``, from :func:`mask_tables`), and
      its ``n``/``n_events`` are shown as the metrics show them.
    - Counts are masked from 1 to ``min_cell - 1``; zero is shown.
    - ``threshold_grid``: :func:`mask_grid_rows`.
    - ``survival_threshold_grid``: a point with 1 to ``min_cell - 1``
      predicted positives or predicted negatives is masked whole.
    - ``risk_distribution``: events and non-events per bin, with
      complementary suppression over the bins of one slice.
    - ``regression_calibration`` and ``calibration_horizon``: the bin count
      ``n`` with complementary suppression over the bins; a masked bin's
      observed value is NaN.
    - ``km_by_risk_group``: an ``at_risk@q<k>`` of 1 to ``min_cell - 1`` is masked
      and so is that quarter's survival at that time.
    """
    import pandas as pd

    label = masked_label(min_cell)
    out = curves.copy().reset_index(drop=True)
    out["masked"] = False
    if not len(out):
        return out

    whole, label_events, label_n = _few_events(out, min_cell)
    _blank(out, whole)
    _label(out, label_events, "n_events", label)
    _label(out, label_n, "n", label)
    out.loc[whole, "masked"] = True
    if hidden_slices:
        keys = _row_keys(out, metric_slice_columns(out))
        hit = pd.Series([key in hidden_slices for key in keys], index=out.index)
        _blank(out, hit)
        out["n"] = out["n"].astype(object)
        out["n_events"] = out["n_events"].astype(object)
        for position in out.index[hit]:
            n_shown, e_shown = hidden_slices[keys[position]]
            out.at[position, "n"] = n_shown
            out.at[position, "n_events"] = e_shown
        out.loc[hit, "masked"] = True

    out = mask_grid_rows(out, min_cell)

    value = pd.to_numeric(out["value"], errors="coerce")
    n = pd.to_numeric(out["n"], errors="coerce")
    survival = (out["curve"] == SURVIVAL_THRESHOLD_GRID) & (out["metric"] == "n_predicted_positive")
    negatives = n - value
    bad = survival & ~out["masked"] & (((value > 0) & (value < min_cell))
                                       | ((negatives > 0) & (negatives < min_cell)))
    _mask_points(out, SURVIVAL_THRESHOLD_GRID, bad, min_cell, ("n_predicted_positive",))

    slice_keys = _slice_columns(out) + ["metric"]
    histogram = (out["curve"] == RISK_DISTRIBUTION) & out["metric"].isin(["events", "nonevents"]) \
        & ~out["masked"]
    if histogram.any():
        # The bins add up to the slice's shown events: protected as far as that allows.
        part = suppress(out[histogram], ["value"], min_cell, slice_keys, strict=False)
        changed = part["value"].map(_is_masked)
        out["value"] = out["value"].astype(object)
        out.loc[part.index[changed], "value"] = part.loc[changed, "value"]
        out.loc[part.index[changed], "masked"] = True

    for curve in (REGRESSION_CALIBRATION, CALIBRATION_HORIZON):
        bins = (out["curve"] == curve) & (out["metric"] == "n") & ~out["masked"]
        if not bins.any():
            continue
        part = suppress(out[bins], ["value"], min_cell, slice_keys, strict=False)
        changed = pd.Series(False, index=out.index)
        changed.loc[part.index] = part["value"].map(_is_masked).to_numpy()
        shown = dict(zip(part.index, part["value"]))
        _mask_points(out, curve, changed, min_cell, ("n",))
        for position in part.index[part["value"].map(_is_masked).to_numpy()]:
            out.at[position, "value"] = shown[position]

    km = out["curve"] == KM_BY_RISK_GROUP
    at_risk = km & out["metric"].astype(str).str.startswith("at_risk@") & ~out["masked"]
    low = at_risk & (value > 0) & (value < min_cell)
    if low.any():
        keys = _x_key(out)
        bad = {(keys[i], out.at[i, "metric"].split("@", 1)[1]) for i in out.index[low]}
        survival_rows = km & out["metric"].astype(str).str.startswith("survival@")
        hit = pd.Series(
            [bool(survival_rows.iat[i]) and (keys[i], str(out.at[i, "metric"]).split("@", 1)[1])
             in bad for i in range(len(out))], index=out.index,
        )
        _blank(out, hit)
        _blank(out, low, ("ci_lo", "ci_hi"))
        _label(out, low, "value", label)
        out.loc[hit | low, "masked"] = True
    return out


def sample_size_table(metrics):
    """Per arm x subgroup x outcome x update: n, events, non-events, >= 100 events, CI widths.

    Read from the masked ``metrics_long`` core rows (its ``n_nonevents``
    when present), so a masked count stays masked and a masked estimate has
    no CI width.
    """
    import pandas as pd

    core = metrics[metrics["tier"] == "core"] if "tier" in metrics.columns else metrics.iloc[0:0]
    keys = [c for c in _slice_columns(core) if c != "curve"]
    rows = []
    for key, part in core.groupby(keys, sort=False, dropna=False):
        key = key if isinstance(key, tuple) else (key,)
        row = dict(zip(keys, key))
        first = part.iloc[0]
        n, events = first["n"], first["n_events"]
        row["n"] = n
        row["n_events"] = events
        e_num = pd.to_numeric(pd.Series([events]), errors="coerce").iat[0]
        n_num = pd.to_numeric(pd.Series([n]), errors="coerce").iat[0]
        if events is None or (not isinstance(events, str) and pd.isna(events)):
            row["n_nonevents"] = None
            row["meets_100_events"] = None
        else:
            if "n_nonevents" in part.columns:
                row["n_nonevents"] = first["n_nonevents"]
            elif isinstance(n, str) or isinstance(events, str):
                row["n_nonevents"] = COMPLEMENT_LABEL
            else:
                row["n_nonevents"] = int(n_num - e_num)
            row["meets_100_events"] = False if isinstance(events, str) else bool(
                e_num >= VALIDATION_EVENTS_TARGET)
        lo = pd.to_numeric(part["ci_lo"], errors="coerce")
        hi = pd.to_numeric(part["ci_hi"], errors="coerce")
        for metric, width in zip(part["metric"], (hi - lo)):
            row[f"ci_width:{metric}"] = None if pd.isna(width) else float(width)
        rows.append(row)
    return pd.DataFrame(rows, columns=None if rows else
                        keys + ["n", "n_events", "n_nonevents", "meets_100_events"])


# ---------------------------------------------------------------------------
# Table 1 and exclusions
# ---------------------------------------------------------------------------

_TABLE1_COUNT_STATS = ("count", "n_persons", "n_ecgs")


def _check_table1_stats(table):
    """Only the summaries describe writes may leave; a minimum or maximum is one person's value."""
    from mival.stages.describe import CONTINUOUS_STATS

    allowed = set(_TABLE1_COUNT_STATS) | {"percent"} | set(CONTINUOUS_STATS)
    unknown = sorted(set(table["stat"].astype(str)) - allowed)
    if unknown:
        raise DisclosureError(f"table1_long has statistics outside the export allowlist: {unknown}")
    return table


def _apply_table1(t1, hidden, min_cell: int, cells: Cells):
    """Table 1 with hidden counts labelled, their percents NaN and thin summaries NaN."""
    import numpy as np
    import pandas as pd

    frame = t1["frame"].copy()
    frame["level"] = frame["level"].astype(str).replace({"nan": ""})
    stat = frame["stat"].astype(str).tolist()
    numbers = pd.to_numeric(frame["value"], errors="coerce").tolist()
    n_values = pd.to_numeric(frame["n"], errors="coerce").tolist()
    hidden_count = set()
    values: List[Any] = []
    for i in range(len(frame)):
        cell = t1["value_cell"][i]
        if cell is not None and cell in hidden:
            hidden_count.add((frame["variable"].iat[i], frame["level"].iat[i], frame["column"].iat[i]))
            values.append(_display(hidden[cell], min_cell))
        elif cell is not None:
            values.append(_int_or(numbers[i]))
        else:
            values.append(numbers[i])
    n_out: List[Any] = []
    n_hidden: List[bool] = []
    for i in range(len(frame)):
        cell = t1["n_cell"][i]
        n_hidden.append(cell is not None and cell in hidden and cells.values[cell] > 0)
        if cell is not None and cell in hidden and cells.values[cell] > 0:
            n_out.append(_display(hidden[cell], min_cell))
        else:
            n_out.append(_int_or(n_values[i]))
        if stat[i] == "percent" and (frame["variable"].iat[i], frame["level"].iat[i],
                                     frame["column"].iat[i]) in hidden_count:
            values[i] = np.nan
        elif stat[i] not in _TABLE1_COUNT_STATS and stat[i] != "percent":
            thin = n_values[i] != n_values[i] or n_values[i] < min_cell
            if thin or (cell is not None and cell in hidden):
                values[i] = np.nan
    for i in _thin_remainder_rows(frame, n_values, stat, min_cell, n_hidden):
        values[i] = np.nan
    frame["value"] = values
    frame["n"] = n_out
    return frame


def _thin_remainder_rows(frame, n_values, stat, min_cell: int, n_hidden=None) -> List[int]:
    """Continuous rows of a variable whose columns add up around a part of 1..min_cell-1 values.

    Columns add up (overall = dev + test; test = its outcome groups + the
    test ECGs without the label), and so do ``n x mean`` and the sums of
    squares behind ``sd``. When one part of such a sum holds 1 to
    ``min_cell - 1`` values, its mean and sd follow from the others even if
    its own row is blank, or it has no row (the unlabelled remainder). The
    audit protects counts, not these summaries, so every continuous summary
    of that variable is blanked in every column: one part shown would bring
    a remainder back through the other relation.

    Blank-or-shown is itself shown, so it may not depend on a count the
    export hides: when any part's (or the total's) ``n`` is hidden after
    masking, the variable is blanked too, whatever the raw ``n`` says.
    """
    by_variable: Dict[Tuple[str, str], Dict[str, float]] = {}
    masked_n: Dict[Tuple[str, str], set] = {}
    rows: Dict[Tuple[str, str], List[int]] = {}
    for i in range(len(frame)):
        if stat[i] in _TABLE1_COUNT_STATS or stat[i] == "percent":
            continue
        key = (str(frame["variable"].iat[i]), str(frame["level"].iat[i]))
        rows.setdefault(key, []).append(i)
        by_variable.setdefault(key, {})[str(frame["column"].iat[i])] = n_values[i]
        if n_hidden is not None and n_hidden[i]:
            masked_n.setdefault(key, set()).add(str(frame["column"].iat[i]))
    blank: List[int] = []
    for key, n_by_column in by_variable.items():
        def n(column):
            value = n_by_column.get(column)
            return 0.0 if value is None or value != value else float(value)

        relations = [(COLUMN_OVERALL, [DEV_SPLIT, TEST_SPLIT])]
        groups = [c for c in n_by_column if c.startswith(f"{TEST_SPLIT}:")]
        if groups:
            relations.append((TEST_SPLIT, groups))
        thin = False
        for total, parts in relations:
            if total not in n_by_column:
                continue
            if masked_n.get(key, set()) & set(parts + [total]):
                thin = True
            sizes = [n(c) for c in parts] + [n(total) - sum(n(c) for c in parts)]
            thin = thin or any(0 < size < min_cell for size in sizes)
        if thin:
            blank.extend(rows[key])
    return blank


def _int_or(value):
    if isinstance(value, float) and value == value and value.is_integer():
        return int(value)
    return value


def mask_table1(table, min_cell: int):
    """``table1_long`` with small counts masked (see :func:`mask_tables`).

    Every count row is a cell; the levels of a categorical variable add up
    to the column's ECG count, a label's levels to its non-missing count
    (whose gap to the ECG count is protected too); overall = dev + test and
    test = its outcome groups for every row. A percent of a hidden count is
    NaN, and so is a continuous summary over fewer than ``min_cell`` values
    or over a hidden count. Only the statistics describe writes are
    accepted: any other ``stat`` (a minimum or maximum is one person's
    value) fails the export.
    """
    return mask_tables(min_cell, table1=table)[0]


#: describe's Table 1 rows that count persons; every other row counts or summarises ECGs
#: (a person with several ECGs is in it once per ECG; ``persons`` gives how many persons).
TABLE1_PER_PERSON = ("persons", "ecgs_per_person")
TABLE1_UNIT = "unit"


def with_table1_unit(table):
    """Table 1 with a ``unit`` column: ``person`` or ``ECG``, what each row counts or summarises.

    With ``ecg_selection.per_person: all`` a person contributes every ECG,
    so age and the other summaries are per ECG, not per person; the column
    says so beside each value (TRIPOD+AI 20b asks for the unit).
    """
    out = table.copy()
    out[TABLE1_UNIT] = ["person" if str(v) in TABLE1_PER_PERSON else "ECG" for v in out["variable"]]
    return out


def table1_wide(table):
    """Rows ``(variable, level, stat[, unit])``, one column per Table 1 column, in table order."""
    import pandas as pd

    rows = ["variable", "level", "stat"] + ([TABLE1_UNIT] if TABLE1_UNIT in table.columns else [])
    columns = list(dict.fromkeys(table["column"].astype(str)))
    index = list(dict.fromkeys(zip(*(table[c] for c in rows))))
    wide = table.set_index(rows + ["column"])["value"].unstack("column")
    wide = wide.reindex(pd.MultiIndex.from_tuples(index, names=rows))
    return wide[[c for c in columns if c in wide.columns]].reset_index()


def _exclusion_cells(cells: Cells, exclusions, tolerate: bool) -> Dict[str, Any]:
    """Cells of ``exclusions_long`` and their constraints.

    One reason's sex levels add up to its overall row, and so do its age-band
    levels; a stage's reasons add up to the stage's total, which the record
    counts before and after the stage give away, so it is a fixed (never
    hidden) cell and only reasons are hidden to protect each other.
    ``tolerate`` marks the reasons :attr:`Cells.tolerated`.
    """
    import pandas as pd

    frame = exclusions.reset_index(drop=True)
    whole, *levels = EXCLUSION_GROUPS
    index = [cells.add(v) for v in pd.to_numeric(frame["n"], errors="coerce")]
    stage = frame["stage"].astype(str).tolist()
    reason = frame["reason_code"].astype(str).tolist()
    group = frame["group"].astype(str).tolist()
    by_stage: Dict[str, List[int]] = {}
    by_reason: Dict[Tuple[str, str], Dict[str, List[int]]] = {}
    for i in range(len(frame)):
        if index[i] is None:
            continue
        if group[i] == whole:
            by_stage.setdefault(stage[i], []).append(index[i])
        by_reason.setdefault((stage[i], reason[i]), {}).setdefault(group[i], []).append(index[i])
    totals: Dict[str, int] = {}
    for name, parts in by_stage.items():
        totals[name] = cells.add(sum(cells.values[p] for p in parts), fixed=True)
        cells.constrain(totals[name], parts)
    for groups in by_reason.values():
        total = (groups.get(whole) or [None])[0]
        for level in levels:
            cells.constrain(total, groups.get(level, []))
    if tolerate:
        cells.tolerated.update(c for c in index if c is not None)
    return {"frame": frame, "index": index, "totals": totals}


def _stard_cells(cells: Cells, stard, totals: Mapping[str, int], metrics: Optional[_Counted],
                 whole) -> Dict[str, List[Optional[int]]]:
    """Cells of the STARD boxes (figure only) and how they meet the tables.

    ``stard`` is evaluate's ``(stages, exclusion lines, side)``: the chain
    ``assessed, after stage 1, ..., after stage k, analysed`` and the dev box.
    """
    stages, _, side = stard
    stage_cells = [cells.add(count, key=("s", "box", i)) for i, (_, count) in enumerate(stages)]
    side_cells = [cells.add(count, key=("s", "side", i)) for i, (_, count) in enumerate(side)]
    chain, analysed = stage_cells[:-1], stage_cells[-1]
    dev = side_cells[0] if side_cells else None
    # Each chain box is the next one plus that stage's exclusions (by construction
    # of the flow); the stage total is the exclusions' fixed cell when it agrees.
    for k in range(len(chain) - 1):
        name = stages[k + 1][0]
        gap = cells.values[chain[k]] - cells.values[chain[k + 1]]
        total = totals.get(name)
        if total is None or cells.values[total] != gap:
            total = cells.add(gap, fixed=True)
        cells.constrain(chain[k], [chain[k + 1], total])
    # The last chain box is what survived: the analysed (test) ECGs and dev.
    if chain:
        cells.constrain(chain[-1], [analysed, dev])
    # analysed and dev are the test and dev ECGs with a prediction: subsets of
    # Table 1's split counts, supersets of every whole metrics slice of the split.
    for cell, split in ((analysed, TEST_SPLIT), (dev, DEV_SPLIT)):
        cells.at_most(cells.get(("t", "ecgs", "", split)), [cell])
        if metrics is None:
            continue
        frame = metrics.frame
        for entry in whole.values():
            i = entry["row"]
            if (str(frame["split"].iat[i]) if "split" in frame.columns else "") != split:
                continue
            if "subgroup" in frame.columns and _axis_of(str(frame["subgroup"].iat[i])) is not None:
                continue
            cells.at_most(cell, sorted(entry["n"]))
    # describe drops nothing: Table 1's overall ECGs may be one of the chain boxes.
    overall = cells.get(("t", "ecgs", "", COLUMN_OVERALL))
    for box in chain:
        if overall is not None and cells.values[box] == cells.values[overall]:
            cells.equal(overall, box)
    return {"stages": stage_cells, "side": side_cells}


def mask_exclusions(exclusions, min_cell: int, strict: bool = True):
    """``exclusions_long`` masked over its additive constraints (see :func:`_exclusion_cells`).

    ``strict`` refuses a small count that only a stage total could protect
    (the export lets it pass only with ``--allow-small-stage-totals``).
    """
    return mask_study(min_cell, exclusions=exclusions, tolerate_exclusions=not strict).exclusions


def stard_display(stages, exclusions, side, masked_exclusions, min_cell: int):
    """STARD boxes and annotations for the bundle figure.

    Exclusion counts are the masked ones. A stage with a single hidden
    reason would give its count away as the difference of the boxes around
    it, so the box before it and every box above are ``withheld`` (with the
    stage totals fixed this happens only when the export allowed a small
    stage total).
    """
    shown = {}
    per_stage: Dict[str, int] = {}
    for row in masked_exclusions[masked_exclusions["group"] == EXCLUSION_GROUPS[0]].itertuples():
        shown[(str(row.stage), str(row.reason_code))] = row.n
        if _is_masked(row.n):
            per_stage[str(row.stage)] = per_stage.get(str(row.stage), 0) + 1
    lines = [(stage, reason, shown.get((stage, reason), n)) for stage, reason, n in exclusions]
    label = masked_label(min_cell)

    def small(count):
        return label if isinstance(count, (int, float)) and 0 < count < min_cell else count

    boxes = [(name, small(count)) for name, count in stages]
    exposed = [i for i, (name, _) in enumerate(boxes) if per_stage.get(name) == 1]
    if exposed:
        last = max(exposed)
        boxes = [(name, WITHHELD) if i < last else (name, count)
                 for i, (name, count) in enumerate(boxes)]
    return boxes, lines, [(name, small(count)) for name, count in side]


# ---------------------------------------------------------------------------
# Paths
# ---------------------------------------------------------------------------


def _root_pattern(root: str):
    root = str(root).strip()
    windows = bool(re.match(r"^[A-Za-z]:", root)) or "\\" in root
    parts = [p for p in re.split(r"[\\/]+", root) if p]
    if not parts:
        return None
    lead = r"[\\/]+" if root[:1] in ("/", "\\") else ""
    body = lead + r"[\\/]+".join(re.escape(p) for p in parts)
    return re.compile(body + r"(?=[\\/]|$|[^A-Za-z0-9._-])", re.IGNORECASE if windows else 0)


def _variants(root) -> List[str]:
    """``root`` as given; a sequence of roots (today's and the recorded ones) as each of them."""
    if root is None:
        return []
    if isinstance(root, (list, tuple, set)):
        return [str(r) for r in root if r is not None and str(r).strip()]
    return [str(root)] if str(root).strip() else []


def _roots(home, runs):
    found = []
    for root, token in ((runs, RUNS_TOKEN), (home, HOME_TOKEN)):
        variants = set()
        for given in _variants(root):
            variants.add(given)
            if not re.match(r"^[A-Za-z]:", given) and given.startswith("/"):
                variants.add(os.path.realpath(given))
        for variant in variants:
            pattern = _root_pattern(variant)
            if pattern is not None:
                found.append((len(variant.rstrip("\\/")), pattern, token))
    return [(p, t) for _, p, t in sorted(found, key=lambda item: -item[0])]


def relativize(obj: Any, home, runs) -> Any:
    """``obj`` with the ``runs`` root written ``$MIVAL_RUNS`` and ``home`` ``$MIVAL_HOME``.

    Works on POSIX and Windows roots alike: a drive path matches in any case
    of letters and with ``\\`` or ``/`` in any mix. The longer root is tried
    first, so runs inside home become ``$MIVAL_RUNS``. A string that starts
    with a root has its remaining separators written ``/``. ``home`` and
    ``runs`` may each be a list of roots: today's, and where a moved run
    recorded them (:class:`RecordedRoots`), which mean the same place.
    """
    roots = _roots(home, runs)

    def fix(text: str) -> str:
        for pattern, token in roots:
            match = pattern.match(text)
            if match:
                return token + text[match.end():].replace("\\", "/")
        for pattern, token in roots:
            text = pattern.sub(lambda _m, t=token: t, text)
        return text

    def walk(value):
        if isinstance(value, dict):
            return {key: walk(item) for key, item in value.items()}
        if isinstance(value, (list, tuple)):
            return [walk(item) for item in value]
        if isinstance(value, str):
            return fix(value)
        return value

    return walk(obj)


def _pure(text: str):
    """``text`` as the kind of path it is spelled as: a drive letter or ``\\`` is Windows."""
    text = str(text)
    if re.match(r"^[A-Za-z]:", text) or "\\" in text:
        return PureWindowsPath(text)
    return PurePosixPath(text)


def _below(path, root) -> Optional[Tuple[str, ...]]:
    """The parts of ``path`` under ``root`` (``()`` for root itself), or None when not under it.

    A Windows path compares in any case of letters, as Windows does.
    """
    if isinstance(path, PureWindowsPath) != isinstance(root, PureWindowsPath):
        return None
    fold = (lambda part: part.lower()) if isinstance(path, PureWindowsPath) else (lambda part: part)
    head = [fold(p) for p in root.parts]
    if [fold(p) for p in path.parts[:len(head)]] != head:
        return None
    return tuple(path.parts[len(head):])


class RecordedRoots:
    """Where a study's runs were when they were recorded, and where they are today.

    ``run_study`` records absolute paths: the run directories in runs.json
    and every stage input, so the input paths of the evaluate manifest too. A
    site that moves its runs afterwards (scratch to the data volume, another
    drive letter) exports with today's ``MIVAL_RUNS``, and a recorded path
    then names a place that is gone, or one that still holds an older copy.

    The recorded runs root is read from runs.json itself: every run directory
    is ``<root>/<study_id>/<stage>/<config_hash>``. The recorded home follows
    when today's runs root sits in today's home at the same relative place
    (``$MIVAL_HOME/runs`` then and now). A recorded path under either is read
    under today's root when it is there, else where it was recorded. Moving
    proves nothing about content: the evaluate inputs the audit counts from
    are still checked against their recorded sha256 (:func:`_evaluate_inputs`).
    """

    def __init__(self, runs_root: PathLike, home: Optional[PathLike], study_id: str,
                 run_dirs: Iterable[PathLike], evidence: Iterable[PathLike] = ()):
        self.runs_root = Path(runs_root)
        self.home = Path(home) if home is not None else None
        #: (recorded root, today's root, token), deepest recorded root first.
        self.pairs: List[Tuple[Any, Path, str]] = []
        today_runs = _pure(str(self.runs_root))
        recorded: List[Any] = []
        for run_dir in run_dirs:
            path = _pure(str(run_dir))
            if len(path.parts) >= 4 and path.parts[-3] == study_id:
                root = path.parents[2]
                if not any(_below(root, seen) == () for seen in recorded):
                    recorded.append(root)
        tail: Optional[Tuple[str, ...]] = None
        if self.home is not None:
            tail = _below(today_runs, _pure(str(self.home)))
        #: Possible old homes: the parent of a recorded runs root that ends the way
        #: today's runs folder does. A guess until :meth:`confirm_home` finds evidence.
        self._candidates: List[Any] = []
        for root in recorded:
            if _below(root, today_runs) == ():
                continue
            self.pairs.append((root, self.runs_root, RUNS_TOKEN))
            if tail and len(root.parts) > len(tail):
                old_home = root.parents[len(tail) - 1]
                fold = (lambda parts: [p.lower() for p in parts]) \
                    if isinstance(root, PureWindowsPath) else list
                if fold(root.parts[-len(tail):]) == fold(tail) and \
                        _below(old_home, _pure(str(self.home))) != () and \
                        len(old_home.parts) >= 2:   # never a filesystem, drive or share root
                    self._candidates.append((old_home, root))
        self.confirm_home(evidence)
        self._sort()

    def _sort(self) -> None:
        self.pairs.sort(key=lambda pair: -len(pair[0].parts))

    def confirm_home(self, evidence: Iterable[PathLike]) -> None:
        """Take a guessed old home only when a recorded path proves it.

        The proof: a recorded path (the inputs the run manifests list) lies
        under the guess, outside the recorded runs root, and the same place
        under today's ``MIVAL_HOME`` exists. A shared trailing folder name
        (``scratch/runs`` and ``data/mival/runs``) proves nothing, and it
        would rewrite unrelated absolute paths into ``$MIVAL_HOME``; those
        stay absolute and the disclosure check refuses them.
        """
        if self.home is None:
            return
        recorded = [_pure(str(p)) for p in evidence]
        for old_home, runs_root in list(self._candidates):
            if any(old == old_home for old, _, kind in self.pairs if kind == HOME_TOKEN):
                continue
            for path in recorded:
                rest = _below(path, old_home)
                if not rest or _below(path, runs_root) is not None:
                    continue
                if self.home.joinpath(*rest).exists():
                    self.pairs.append((old_home, self.home, HOME_TOKEN))
                    break
        self._sort()

    def path(self, recorded: PathLike) -> Path:
        """``recorded`` under today's root when it is there, else as recorded."""
        text = str(recorded)
        path = _pure(text)
        for old, new, _ in self.pairs:
            rest = _below(path, old)
            if rest is not None:
                moved = new.joinpath(*rest)
                if moved.exists():
                    return moved
        return Path(text)

    def roots(self, token: str) -> List[str]:
        """The recorded roots that mean ``$MIVAL_RUNS`` (or ``$MIVAL_HOME``) today."""
        return [str(old) for old, _, kind in self.pairs if kind == token]


# ---------------------------------------------------------------------------
# Disclosure check
# ---------------------------------------------------------------------------


def _forbidden_name(name: Any) -> bool:
    text = str(name).strip()
    return text.lower() in FORBIDDEN_COLUMNS or bool(_DATE_NAME.search(text))


def _check_value(value: Any, where: str, allow_date: bool = False) -> None:
    if not isinstance(value, str):
        return
    if _ABSOLUTE.search(value) or _EMBEDDED_ABSOLUTE.search(value):
        raise DisclosureError(f"{where}: absolute path in {value!r}")
    if not allow_date and _DATE_VALUE.search(value):
        raise DisclosureError(f"{where}: date in {value!r}")


def _check_names(names: Iterable[Any], where: str) -> None:
    for name in names:
        if _forbidden_name(name):
            raise DisclosureError(f"{where}: forbidden column {name!r}")
        _check_value(str(name), where)


def _check_csv(path: Path) -> None:
    import pandas as pd

    try:
        frame = pd.read_csv(path, dtype=str, keep_default_na=False)
    except pd.errors.EmptyDataError:
        return
    _check_names(frame.columns, str(path))
    for column in frame.columns:
        for value in frame[column]:
            _check_value(value, f"{path} [{column}]")


def _check_xlsx(path: Path) -> None:
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True)
    try:
        for sheet in book.worksheets:
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                where = f"{path} [{sheet.title}]"
                if index == 0:
                    _check_names([c for c in row if c is not None], where)
                for value in row:
                    _check_value(value if isinstance(value, str) else None, where)
    finally:
        book.close()


def _check_json(path: Path) -> None:
    body = json.loads(path.read_text(encoding="utf-8"))
    is_export = path.name == Path(EXPORT_MANIFEST).name

    def walk(value, where, top=False):
        if isinstance(value, dict):
            for key, item in value.items():
                if _forbidden_name(key):
                    raise DisclosureError(f"{path}: forbidden key {key!r} at {where}")
                _check_value(str(key), f"{path} {where}")
                if top and is_export and key == GENERATED_AT:
                    _check_value(item, f"{path} {where}{key}", allow_date=True)
                    continue
                walk(item, f"{where}{key}.")
        elif isinstance(value, list):
            for item in value:
                walk(item, where)
        else:
            _check_value(value, f"{path} {where}")

    walk(body, "", top=True)


def _check_markdown(path: Path) -> None:
    for number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), start=1):
        where = f"{path}:{number}"
        found = _MD_FORBIDDEN.search(line)
        if found:
            raise DisclosureError(f"{where}: forbidden column {found.group(1)!r}")
        _check_value(line, where)


_CHECKERS = {".csv": _check_csv, ".xlsx": _check_xlsx, ".json": _check_json,
             ".md": _check_markdown}
#: Written only by the figure functions from masked tables; no text to read.
_UNCHECKED = (".png",)


def check_no_patient_level(paths: Iterable[PathLike]) -> None:
    """Raise :class:`DisclosureError` when a file could carry a record-level value.

    csv, xlsx, json and md files are read back. A record identifier, file
    path or individual-date column (or JSON key), an absolute path (POSIX,
    drive letter in either slash, or UNC) anywhere in a value, or a date
    ``YYYY-MM-DD`` fails; the export manifest's own ``generated_at`` is the
    one date allowed. A file of any other type than these and PNG fails too,
    since it cannot be checked.
    """
    for raw in paths:
        path = Path(raw)
        suffix = path.suffix.lower()
        if suffix in _UNCHECKED:
            continue
        checker = _CHECKERS.get(suffix)
        if checker is None:
            raise DisclosureError(f"{path}: file type {suffix or '(none)'} cannot be checked")
        checker(path)


# ---------------------------------------------------------------------------
# The bundle
# ---------------------------------------------------------------------------


#: What a stage summary may export, per summary file: ``{key: sub-keys}``.
#: Every count (record totals, splits, the gate, positives, exclusions, the
#: label delays) is left out: they overlap Table 1, the exclusions and the
#: STARD boxes, so a summary count could undo their masking. Only the label
#: value's quantiles are kept, and only as floats.
SUMMARY_ALLOWLIST: Mapping[str, Mapping[str, Tuple[str, ...]]] = {
    "retrieve": {"label_value_quantiles": ("p5", "p25", "p50", "p75", "p95")},
    "profile": {"label_value_quantiles": ("p5", "p25", "p50", "p75", "p95")},
}


def provenance_summary(stage: str, raw: Mapping[str, Any]) -> Dict[str, Any]:
    """The allowlisted, count-free part of one stage summary (see :data:`SUMMARY_ALLOWLIST`).

    A key not listed is dropped, and so is any value that is not a float
    (an integer is a count, whatever its key).
    """
    out: Dict[str, Any] = {}
    for key, inner in SUMMARY_ALLOWLIST.get(stage, {}).items():
        body = raw.get(key) if isinstance(raw, Mapping) else None
        if not isinstance(body, Mapping):
            continue
        kept = {k: body[k] for k in inner if isinstance(body.get(k), float)}
        if kept:
            out[key] = kept
    return out


def _read_json(path: Path) -> Any:
    return json.loads(Path(path).read_text(encoding="utf-8"))


def _write_json(path: Path, body: Any) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(body, indent=2, ensure_ascii=False, default=str) + "\n",
                    encoding="utf-8")
    return path


def _write_csv(path: Path, frame) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    frame.to_csv(path, index=False)
    return path


def _has_openpyxl() -> bool:
    import importlib.util

    return importlib.util.find_spec("openpyxl") is not None


def _has_matplotlib() -> bool:
    import importlib.util

    return importlib.util.find_spec("matplotlib") is not None


def _sheet(name: str, used: set) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "-", name)[:31]
    candidate, k = base, 1
    while candidate in used:
        k += 1
        candidate = f"{base[:28]}~{k}"
    used.add(candidate)
    return candidate


def _write_xlsx(path: Path, sheets: Sequence[Tuple[str, Any]]) -> Path:
    import pandas as pd

    path.parent.mkdir(parents=True, exist_ok=True)
    used: set = set()
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, frame in sheets:
            frame.to_excel(writer, sheet_name=_sheet(name, used), index=False)
    return path


def _cuts_record(value: Any) -> Any:
    """The regression cuts as ``[{value, op}]``, or per label_def when the spec maps them."""
    def items(cuts):
        return [{"value": c.value, "op": c.op} for c in criteria.parse_cuts(cuts, "regression_cuts")]

    if isinstance(value, Mapping) and not criteria._is_cut_item(value):
        return {str(label): items(cuts) for label, cuts in value.items()}
    return items(value)


def study_criteria(study) -> Dict[str, Any]:
    """The thresholds a study's export ran under, as recorded in the export manifest."""
    from mival.stages.evaluate import SUBGROUP_SUPPRESS_MIN_EVENTS
    from mival.stages.retrieve import DEFAULT_FOLLOW_UP_DAYS

    describe = study.stage_spec("describe")
    source = study.stage_spec("retrieve").get("label_source") or {}
    follow_up: Dict[str, Any] = {}
    if source.get("kind") == "death_within":
        # The death definition the site ran under: how long after the last
        # visit a death still counts as recorded (later follow-up is censored).
        stated = source.get("follow_up_after_last_visit_days")
        follow_up = {
            "follow_up_after_last_visit_days": int(DEFAULT_FOLLOW_UP_DAYS if stated is None else stated),
            "follow_up_after_last_visit_days_from": "study.yaml" if stated is not None
            else "engine default",
        }
    evaluate = study.stage_spec("evaluate")
    bands = criteria.age_bands(describe)
    grid = criteria.net_benefit_grid(evaluate)
    spec_grid = dict(criteria.DEFAULT_NET_BENEFIT_GRID)
    spec_grid.update(evaluate.get("net_benefit_grid") or {})
    return {
        "age_bands": [float(b) for b in bands],
        "age_band_labels": list(criteria.age_band_labels(bands)),
        # Used by regression arms only. There is no engine default: without the
        # key a regression arm has no cut metrics, and this records that.
        "regression_cuts": _cuts_record(evaluate["regression_cuts"])
        if "regression_cuts" in evaluate else [],
        "regression_cuts_from": "study.yaml" if "regression_cuts" in evaluate
        else "not set: no cut metrics",
        "net_benefit_grid": {"start": float(spec_grid["start"]), "stop": float(spec_grid["stop"]),
                             "step": float(spec_grid["step"]), "n_points": int(grid.size)},
        "min_cell": criteria.min_cell(study.stage_spec("export")),
        "min_events": int(evaluate.get("min_events", SUBGROUP_SUPPRESS_MIN_EVENTS)),
        **follow_up,
    }


def _load_runs_file(runs_root: Path, study_id: str, home: Optional[PathLike] = None
                    ) -> Tuple[Path, Dict[str, List[Path]], "RecordedRoots"]:
    """runs.json of ``study_id``, its run directories read under today's roots."""
    path = Path(runs_root) / study_id / RUNS_FILE
    if not path.is_file():
        raise ValueError(f"study {study_id}: no runs.json at {path}; run the study first")
    raw = _read_json(path)
    roots = RecordedRoots(runs_root, home, study_id, [d for dirs in raw.values() for d in dirs])
    seen: List[str] = []
    for dirs in raw.values():
        for d in dirs:
            try:
                listed = _read_json(roots.path(d) / "manifest.json").get("input_artifacts") or ()
            except (OSError, ValueError, AttributeError):
                continue
            seen.extend(str(ref.get("path")) for ref in listed
                        if isinstance(ref, dict) and ref.get("path"))
    roots.confirm_home(seen)
    runs = {str(stage): [roots.path(d) for d in dirs] for stage, dirs in raw.items()}
    for stage in REQUIRED_STAGES:
        if not runs.get(stage):
            raise ValueError(
                f"study {study_id}: runs.json lists no {stage} run; the bundle needs this "
                f"invocation's {', '.join(REQUIRED_STAGES)} output"
            )
    return path, runs, roots


def _artifact(run_dir: Path, name: str) -> Path:
    return Path(run_dir) / "artifacts" / name


def _concat_tables(run_dirs: Sequence[Path], name: str):
    import pandas as pd

    from mival.pipeline.tables import read_table

    parts = [read_table(_artifact(d, name)) for d in run_dirs if _artifact(d, name).is_file()]
    if not parts:
        return None
    return pd.concat(parts, ignore_index=True) if len(parts) > 1 else parts[0]


def _single(run_dirs: Sequence[Path], stage: str, study_id: str) -> Path:
    if len(run_dirs) != 1:
        raise ValueError(f"study {study_id}: runs.json lists {len(run_dirs)} {stage} runs; expected 1")
    return run_dirs[0]


class _Bundle:
    """The files of one bundle as they are written into a staging directory."""

    def __init__(self, root: Path, min_cell: int, home, runs):
        self.root = root
        self.min_cell = min_cell
        self.home = home
        self.runs = runs
        self.notes: List[str] = []
        #: ``--allow-small-stage-totals``: proceed, and record the stages.
        self.allow_small_stage_totals = False
        self.small_stage_totals: Dict[str, Dict[str, int]] = {}
        self.figure_notes: List[str] = []

    def path(self, relative: str) -> Path:
        if not any(re.fullmatch(p, relative) for p in ALLOWLIST):
            raise DisclosureError(f"{relative} is not on the export allowlist")
        target = self.root / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        return target

    def files(self) -> List[str]:
        return sorted(p.relative_to(self.root).as_posix() for p in self.root.rglob("*") if p.is_file())


def _draw(bundle: _Bundle, name: str, function, frame, **kwargs) -> None:
    target = bundle.path(f"figures/{_safe(name)}.png")
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        function(frame, target, **kwargs)
    for warning in caught:
        bundle.figure_notes.append(f"{target.name}: {warning.message}")


def _safe(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]", "-", name)


_CURVE_FIGURES = (
    (THRESHOLD_GRID, "roc_pr", "grid_roc_pr"),
    (THRESHOLD_GRID, "decision_curve", "grid_decision_curve"),
    (CALIBRATION, "calibration", "calibration_band"),
    (REGRESSION_CALIBRATION, "regression_calibration", "regression_calibration_plot"),
    (KM_BY_RISK_GROUP, "km", "km_plot"),
    (SURVIVAL_THRESHOLD_GRID, "survival_decision_curve", "survival_decision_curve"),
)
#: Subgroup forest plots: the lead core metric of each kind of arm.
_FOREST_STEMS = ("auroc", "mae", "uno_c")


def _figures(bundle: _Bundle, study_id: str, curves, metrics, stard) -> None:
    from mival import figures
    from mival.stages.evaluate import SUBGROUP_ALL, UPDATE_NONE

    if stard is not None:
        boxes, lines, side = stard
        target = bundle.path(f"figures/{study_id}_stard_flow.png")
        figures.stard_from_counts(boxes, lines, target, side=side)

    if curves is not None and len(curves):
        shipped = curves[(curves["update"] == UPDATE_NONE) & (curves["subgroup"] == SUBGROUP_ALL)]
        for curve, name, function in _CURVE_FIGURES:
            rows = shipped[shipped["curve"] == curve]
            outcomes = sorted(set(rows["outcome"].astype(str)))
            for outcome in outcomes:
                label = f"{study_id}_{name}" if len(outcomes) == 1 else f"{study_id}_{outcome}_{name}"
                _draw(bundle, label, getattr(figures, function), curves, outcome=outcome)

    if metrics is not None and len(metrics):
        test = metrics[(metrics["split"].astype(str) == TEST_SPLIT)
                       & (metrics["update"] == UPDATE_NONE)]
        has_subgroups = test[test["subgroup"] != SUBGROUP_ALL]
        stems = test["metric"].astype(str).str.split("@", n=1).str[0]
        for metric in sorted(set(test["metric"][stems.isin(_FOREST_STEMS)].astype(str))):
            if not (has_subgroups["metric"] == metric).any():
                continue
            rows = test[test["metric"] == metric]
            for outcome in sorted(set(rows["outcome"].astype(str))):
                many = rows["outcome"].astype(str).nunique() > 1
                label = f"{study_id}_forest_{metric}" + (f"_{outcome}" if many else "")
                _draw(bundle, label, lambda frame, path, **kw: figures.forest_from_metrics(
                    frame, metric, path, **kw), rows, outcome=outcome)


def _stard(study, runs_root: Path, runs_json: Path, evaluate_dirs):
    """The STARD flow of this invocation, through evaluate's own counting (unmasked).

    ``(stages, exclusion lines, side)``, or None when the evaluate manifest
    has no STARD counts.
    """
    from mival.stages.evaluate import RUNS_INPUT, STARD_COUNT_ANALYSED, STARD_COUNT_DEV, _stard_counts

    counts = _read_json(Path(evaluate_dirs[0]) / "manifest.json").get("counts") or {}
    if STARD_COUNT_ANALYSED not in counts or STARD_COUNT_DEV not in counts:
        return None
    ctx = SimpleNamespace(
        layout=SimpleNamespace(runs_root=Path(runs_root)),
        study_id=study.study_id,
        inputs={RUNS_INPUT: runs_json},
    )
    return _stard_counts(ctx, counts[STARD_COUNT_ANALYSED], counts[STARD_COUNT_DEV])


def _small_stage_totals(exclusions, min_cell) -> Dict[str, int]:
    """Stages whose whole exclusion count is 1 to min_cell - 1.

    The count is the difference of the record counts before and after the
    stage (STARD boxes, Table 1, metrics ``n``), so masking its reasons does
    not hide it.
    """
    import pandas as pd

    whole = exclusions[exclusions["group"] == EXCLUSION_GROUPS[0]]
    totals = pd.to_numeric(whole["n"], errors="coerce").groupby(whole["stage"]).sum()
    return {str(stage): int(total) for stage, total in totals.items() if 0 < total < min_cell}


def provenance_manifest(raw: Mapping[str, Any], home, runs) -> Dict[str, Any]:
    """A stage manifest reduced to an explicit allowlist of structured fields.

    Free text (``warnings``, ``errors``), timestamps and stage counts are left
    out: a warning may name a case or carry a small count, a timestamp is a
    date, and in/out counts give per-stage exclusion totals. Input artifacts
    keep their (relativized) paths only.
    """
    def pick(mapping, keys):
        mapping = mapping if isinstance(mapping, Mapping) else {}
        return {key: mapping.get(key) for key in keys}

    contamination = raw.get("contamination") if isinstance(raw.get("contamination"), Mapping) else {}
    body = {
        "run_id": raw.get("run_id"),
        "stage": raw.get("stage"),
        "config_hash": raw.get("config_hash"),
        "seed": raw.get("seed"),
        "code": pick(raw.get("code"), ("mival_version", "source_sha256", "container_digest")),
        "git": pick(raw.get("git"), ("commit", "dirty")),
        "env": pick(raw.get("env"), ("python", "platform", "libraries", "packages", "cpu",
                                     "cuda", "gpu", "container_digest")),
        "compute": pick(raw.get("compute"), ("gpu_hours", "wall_time_s")),
        "contamination": {
            "flag": bool(contamination.get("flag", False)),
            "overlapping_corpora": [str(c) for c in contamination.get("overlapping_corpora") or ()],
        },
        "input_artifacts": [{"path": ref.get("path")} for ref in raw.get("input_artifacts") or ()
                            if isinstance(ref, Mapping)],
    }
    return relativize(body, home, runs)


def repository_root() -> Path:
    """The checkout this package runs from (``<repo>/src/mival``), where every stage runs."""
    return Path(__file__).resolve().parents[2]


def registry_dir(registry: PathLike, repo: Optional[PathLike] = None) -> Path:
    """``models.registry`` as the stages read it: absolute as given, else from the repository root.

    The site scripts run every stage with the repository as working
    directory (``run_models._registry``, ``check_bundle``), so a relative
    registry such as ``registry/site/models`` means the repository's. The
    export runs from ``$MIVAL_HOME`` instead; reading the registry against
    that, or against the study directory, finds no card at all.
    """
    path = Path(str(registry))
    return path if path.is_absolute() else Path(repo or repository_root()) / path


def _card_facts(study, home: Path, repo: Optional[PathLike] = None) -> List[Dict[str, Any]]:
    """Arms of the models stage with what their cards say (pretraining data, output)."""
    models = study.stage_spec("models")
    registry = models.get("registry")
    cards: Dict[str, Any] = {}
    base = registry_dir(registry, repo) if registry else None
    if base is not None and base.is_dir():
        for path in sorted(base.glob("*.json")):
            try:
                body = _read_json(path)
            except (OSError, ValueError):
                continue
            if isinstance(body, dict) and "model_id" in body:
                cards[str(body["model_id"])] = body
    arms = []
    for arm in models.get("arms") or ():
        card = cards.get(str(arm.get("model_id")), {})
        extension = card.get("x-mival", {}) if isinstance(card, dict) else {}
        arms.append({
            "model_id": str(arm.get("model_id")),
            "training_mode": str(arm.get("training_mode")),
            "label_def": str(arm.get("label_def")),
            "pretraining_corpora": list(extension.get("pretraining_corpora") or ()),
            "output": dict(extension.get("output") or {}),
        })
    return arms


def runtime_overrides(study, models_dirs: Sequence[Path]) -> List[Dict[str, Any]]:
    """Per arm, the platform runtime override its models run was built with.

    Read from each models run's ``train_log.jsonl`` (``model_card.runtime_override``,
    written only when one was applied), never from the card as it reads at
    export time: the export may run with a different platform setting than the
    models stage did. ``recorded`` is False when no models run of this export
    has a train_log entry for the arm, so None there means "not known", not
    "none applied". Only the override's name and platform leave, and the name
    of the state-dict transform it declared (``state_dict_transform``, else
    None); its list of changes stays in the train_log.
    """
    found: Dict[Tuple[str, str, str], Optional[Dict[str, Any]]] = {}
    for run_dir in models_dirs:
        path = _artifact(run_dir, "train_log.jsonl")
        if not path.is_file():
            continue
        for line in path.read_text(encoding="utf-8").splitlines():
            if not line.strip():
                continue
            try:
                entry = json.loads(line)
            except ValueError:
                continue
            key = (str(entry.get("model_id")), str(entry.get("training_mode")),
                   str(entry.get("label_def")))
            override = (entry.get("model_card") or {}).get("runtime_override")
            if override:
                transform = override.get("state_dict_transform")
                found[key] = {"name": str(override.get("name")),
                              "platform": str(override.get("platform")),
                              "state_dict_transform": None if transform is None else str(transform)}
            else:
                found.setdefault(key, None)
    arms = []
    for arm in study.stage_spec("models").get("arms") or ():
        key = (str(arm.get("model_id")), str(arm.get("training_mode")), str(arm.get("label_def")))
        override = found.get(key)
        arms.append({"model_id": key[0], "training_mode": key[1], "label_def": key[2],
                     "recorded": key in found,
                     "runtime_override": None if override is None else override["name"],
                     "runtime_platform": None if override is None else override["platform"],
                     "state_dict_transform": None if override is None else override["state_dict_transform"]})
    return arms


def _prespecified(study) -> Optional[List[str]]:
    """``export.prespecified_subgroups``: the subgroups the protocol named in advance, or None.

    A protocol fact, so the study states it; without the key the Methods
    draft leaves it to the authors instead of guessing.
    """
    named = study.stage_spec("export").get("prespecified_subgroups")
    return None if named is None else [str(s) for s in named]


def _ecg_years(table) -> Optional[Tuple[int, int]]:
    import pandas as pd

    rows = table[(table["variable"] == "ecg_year") & (table["column"] == COLUMN_OVERALL)
                 & (table["stat"] == "count")]
    years = [int(float(level)) for level, value in zip(rows["level"], rows["value"])
             if str(level).replace(".", "", 1).isdigit()
             and (isinstance(value, str) or (pd.notna(value) and float(value) > 0))]
    return (min(years), max(years)) if years else None


def _export_study(bundle: _Bundle, study, runs_root: Path, home: Path) -> Dict[str, Any]:
    import pandas as pd

    from mival.stages.describe import EXCLUSIONS_LONG, TABLE1_LONG
    from mival.stages.evaluate import (
        BOOTSTRAP_REPLICATES,
        COMPARISON_COLUMNS,
        COMPARISONS,
        CURVES_LONG,
        METRICS_LONG,
        OUTCOME_DIAGNOSTIC,
    )

    sid = study.study_id
    min_cell = bundle.min_cell
    runs_json, runs, roots = _load_runs_file(runs_root, sid, home)
    describe_dir = _single(runs["describe"], "describe", sid)

    table = pd.read_csv(_artifact(describe_dir, TABLE1_LONG), keep_default_na=False)
    evaluate_dirs = runs["evaluate"]
    metrics = _concat_tables(evaluate_dirs, METRICS_LONG)
    if metrics is None:
        raise ValueError(f"study {sid}: evaluate run has no {METRICS_LONG}")
    comparisons = _concat_tables(evaluate_dirs, COMPARISONS)
    if comparisons is None:
        comparisons = pd.DataFrame(columns=list(COMPARISON_COLUMNS))
    outcomes = study.stage_spec("evaluate").get("outcomes") or {OUTCOME_DIAGNOSTIC: None}
    exclusions = pd.read_csv(_artifact(describe_dir, EXCLUSIONS_LONG), keep_default_na=False)
    small_stages = _small_stage_totals(exclusions, min_cell)
    if small_stages:
        listed = ", ".join(f"{stage} ({total})" for stage, total in small_stages.items())
        if not bundle.allow_small_stage_totals:
            raise DisclosureError(
                f"study {sid}: the stages {listed} excluded 1 to {min_cell - 1} records in "
                "total; that count can be recovered from the record counts before and after "
                "the stage. Export refused; rerun with --allow-small-stage-totals only if the "
                "site has decided this count may leave."
            )
        bundle.small_stage_totals[sid] = dict(small_stages)
        bundle.notes.append(
            f"{sid}: `--allow-small-stage-totals`로 반출한다. 제외 수 합계가 1–{min_cell - 1}건인 "
            f"단계: {listed}. 단계 앞뒤 인원의 차로 이 수를 알 수 있다."
        )

    # Read here from evaluate's own inputs; never written to the bundle.
    scored = scored_counts(evaluate_dirs, study.stage_spec("evaluate"), sid, roots)

    # One constraint system for every count this study exports.
    flow = _stard(study, runs_root, runs_json, evaluate_dirs)
    masked = mask_study(
        min_cell, table1=table, metrics=metrics, comparisons=comparisons, exclusions=exclusions,
        stard=flow, outcome_variable=study.stage_spec("describe").get("outcome"),
        primary_outcome=str(next(iter(outcomes))),
        tolerate_exclusions=bundle.allow_small_stage_totals,
        evaluate_spec=study.stage_spec("evaluate"), scored=scored,
    )
    table1, metrics, comparisons, hidden = (masked.table1, masked.metrics, masked.comparisons,
                                            masked.slices)
    masked_exclusions = masked.exclusions
    table1 = with_table1_unit(table1)
    _write_csv(bundle.path(f"table1/{sid}.csv"), table1)
    if _has_openpyxl():
        _write_xlsx(bundle.path(f"table1/{sid}.xlsx"),
                    [("table1", table1_wide(table1)), ("long", table1)])
    _write_csv(bundle.path(f"exclusions/{sid}.csv"), masked_exclusions)

    _write_csv(bundle.path(f"metrics/{sid}_metrics.csv"), metrics)
    _write_csv(bundle.path(f"metrics/{sid}_comparisons.csv"), comparisons)
    sizes = sample_size_table(metrics)
    _write_csv(bundle.path(f"metrics/{sid}_sample_size.csv"), sizes)

    curves = _concat_tables(evaluate_dirs, CURVES_LONG)
    if curves is not None:
        curves = mask_curves(curves, min_cell, hidden)
        _write_csv(bundle.path(f"curves/{sid}_curves.csv"), curves)

    if _has_matplotlib():
        stard = None
        if flow is not None:
            stages, side = masked.stard
            stard = stard_display(stages, flow[1], side, masked_exclusions, min_cell)
        if stard is None:
            bundle.notes.append(f"{sid}: evaluate manifest에 STARD 건수가 없어 STARD 그림을 그리지 않았다.")
        _figures(bundle, sid, curves, metrics, stard)

    stages: Dict[str, List[Dict[str, Any]]] = {}
    for stage, run_dirs in runs.items():
        if stage not in MANIFEST_STAGES:
            continue
        manifests = []
        for run_dir in run_dirs:
            path = Path(run_dir) / "manifest.json"
            if not path.is_file():
                raise ValueError(f"study {sid}: {stage} run {run_dir} has no manifest.json")
            raw = _read_json(path)
            stages.setdefault(stage, []).append({
                "config_hash": raw.get("config_hash"),
                "source_sha256": (raw.get("code") or {}).get("source_sha256"),
            })
            manifests.append(provenance_manifest(
                raw, [bundle.home] + roots.roots(HOME_TOKEN), [bundle.runs] + roots.roots(RUNS_TOKEN)))
        body = manifests[0] if len(manifests) == 1 else {"runs": manifests}
        _write_json(bundle.path(f"provenance/{sid}_{stage}_manifest.json"), body)
        summaries = [_artifact(d, SUMMARY_FILES[stage]) for d in run_dirs] \
            if stage in SUMMARY_FILES else []
        summaries = [_read_json(p) for p in summaries if p.is_file()]
        if summaries:
            kept = [provenance_summary(stage, s) for s in summaries]
            _write_json(bundle.path(f"provenance/{sid}_{stage}_summary.json"),
                        kept[0] if len(kept) == 1 else {"runs": kept})

    evaluate_spec = study.stage_spec("evaluate")
    facts = {
        "study_id": sid,
        "site": study.site,
        "seed": study.seed,
        "ecg_years": _ecg_years(table1),
        "retrieve": study.stage_spec("retrieve"),
        "profile": study.stage_spec("profile"),
        "arms": _card_facts(study, home),
        "subgroups": [str(s) for s in evaluate_spec.get("subgroups") or ()],
        "prespecified_subgroups": _prespecified(study),
        "criteria": study_criteria(study),
        "bootstrap_replicates": int(evaluate_spec.get("bootstrap_replicates", BOOTSTRAP_REPLICATES)),
        "min_cell": min_cell,
    }
    return {"stages": stages, "facts": facts, "metrics": metrics, "comparisons": comparisons,
            "sample_size": sizes, "arms": runtime_overrides(study, runs.get("models") or [])}


_CONTENTS = (
    ("table1/", "Table 1 (긴 형식 csv, xlsx는 넓은 표와 긴 표)"),
    ("metrics/", "지표(metrics), arm 쌍 비교(comparisons), 표본 크기(sample_size)"),
    ("curves/", "고정 격자 곡선: 임계값 격자, calibration, 위험 분포, 회귀 calibration, KM"),
    ("figures/", "가림이 끝난 표에서 다시 그린 그림"),
    ("exclusions/", "단계 × 사유 제외 건수, 전체·sex별·age_band별"),
    ("provenance/", "단계별 manifest(경로는 $MIVAL_HOME/$MIVAL_RUNS 기준), retrieve·profile 요약(건수 없이 "
                    "라벨 값 분위수만), 반출 기준"),
    ("reporting/", "TRIPOD+AI 체크리스트 초안, Methods 초안"),
)


def _readme(bundle: _Bundle, site: str, by_study: Mapping[str, Mapping[str, Any]],
            min_cells: Mapping[str, int]) -> str:
    from mival.reporting import cut_text

    label = masked_label(bundle.min_cell)
    lines = [
        f"# 결과 묶음: {site}",
        "",
        "환자·ECG 단위 값이 없는 집계 결과만 들어 있다. 반출 전에 아래 검토 포인트를 확인한다.",
        "",
        "## 안에 든 것",
        "",
    ]
    lines += [f"- `{folder}` {text}" for folder, text in _CONTENTS]
    lines += ["", "파일:", ""] + [f"- `{name}`" for name in bundle.files()] + ["", "## 반출 검토 포인트", ""]
    lines += [
        f"1. 표에 적힌 건수가 {bundle.min_cell} 이상이거나 `{label}`, 또는 0인지: Table 1, "
        "exclusions, metrics·comparisons의 n·n_events와 tp/fp/tn/fn, curves의 tp/fp/tn/fn, "
        "n_predicted_positive, 위험 구간 events/nonevents, 보정 구간 n, at_risk. 아래 "
        "'가림으로 막지 않는 것'도 읽는다.",
        "2. 환자·ECG 식별자, 파일 경로, 개별 날짜가 어느 파일에도 없는지. 묶음을 만들 때 자동으로 "
        "점검했지만 열어서 한 번 더 본다.",
        "3. provenance의 경로가 `$MIVAL_HOME`, `$MIVAL_RUNS`로 시작하는지.",
        "4. figures는 가림이 끝난 표에서 그렸다. 개별 점을 그리는 그림(회귀 산점도)은 넣지 않았다.",
        "5. Table 1의 장비 제조사 이름이 기관을 드러내도 되는지.",
    ]
    for note in bundle.notes:
        lines.append(f"- {note}")
    lines += [
        "",
        "## 가림 규칙",
        "",
        f"- `{label}` = 반출하지 않은 칸이다. 1–{bundle.min_cell - 1}건인 모든 칸과, 그 칸을 "
        "보호하려고 함께 가린 칸이 모두 같은 표시를 쓴다(어느 칸이 작은지 표시로 알 수 없게). "
        "0은 그대로 둔다.",
        "- 더해서 합계가 되는 칸들을 한 제약으로 본다: subgroup 수준들과 전체 행(n, 사건, 비사건 "
        "각각), n = 사건 + 비사건, dev fold들, 같은 기록을 센 다른 arm·재보정 행, Table 1의 "
        "수준들과 열의 ECG 수, overall = dev + test, test = 결과군 열들, 제외 사유의 sex·age_band "
        "수준과 그 사유의 합계, 단계의 사유들과 단계 합계(단계 합계는 앞뒤 인원의 차로 알려져 있으므로 "
        "가리지 않는다), STARD 상자들(각 상자 = 다음 상자 + 그 단계 제외 수, 마지막 상자 = "
        "analysed + dev).",
        "- 파일 사이의 관계도 같은 제약 체계에 넣는다(study 하나에 체계 하나). metrics slice의 "
        "기록은 Table 1의 같은 split·수준 기록의 일부다. 어느 표에도 나오지 않는 수도 칸으로 넣는다: "
        "라벨(outcome × label_def)마다 그 split·수준의 라벨 1·0·없음 ECG 수(describe.outcome 열로 "
        "채점한 라벨은 Table 1이 보여 주는 split 전체와 test 수준별 결과군 열의 건수를 쓴다. 같은 "
        "outcome 이름이라도 다른 열로 채점한 arm, 예컨대 생존 arm의 사건은 그 건수와 묶지 않는다). arm마다 그 "
        "split·수준에서 예측한 ECG 수(arm이 그 split·수준의 Table 1 ECG를 모두 예측했으면 읽는 사람이 "
        "그 사실을 안다고 본다. 같은 기관에서는 모든 arm이 Table 1의 ECG를 다 예측했다고 짐작할 수 있다). "
        "관계는 Table 1 건수 = 라벨 1 + 0 + 없음, 수준들의 합 = split "
        "전체, slice n ≤ arm이 예측한 ECG ≤ Table 1 건수, arm이 예측한 ECG ≤ 라벨 없음 + n, "
        "사건 ≤ 라벨 1, 비사건 ≤ 라벨 0이다. 라벨 1·0·없음 수가 1–"
        f"{bundle.min_cell - 1}일 수 있다고 읽히면 다른 칸처럼 보호한다. comparisons는 두 arm slice의 일부이고, STARD의 "
        "analysed·dev는 Table 1 test·dev ECG 수의 일부이자 그 split 모든 전체 slice n 이상이다. "
        "이 부등식들은 건수가 같으면 같은 기록으로 보고 등식으로 쓴다(한쪽에서 가리면 다른 "
        "쪽에서도 가린다). sample_size는 metrics의 칸을 그대로 보여 준다.",
        "- 합계가 보이는 제약에서는 가린 칸이 없거나, 둘 이상이면서 그 합이 "
        f"{bundle.min_cell} 이상이어야 한다. 합계를 가렸으면 부분이 모두 보여서는 안 된다. 어긋나면 "
        "가장 작은 양수 칸을 더 가리고 다시 본다(보완 가림).",
        "- 여러 합계를 함께 풀어도 작은 칸이 드러나지 않는지 선형계획으로 확인한다. 보이는 칸과 합계를 "
        "알고 가린 칸은 1 이상이라고 볼 때, 작은 칸마다 가능한 범위의 최댓값이 "
        f"{bundle.min_cell} 이상이고 폭이 {bundle.min_cell - 1} 이상이어야 한다. 아니면 칸을 더 가리고, "
        "더 가릴 칸이 없으면 반출을 멈춘다.",
        "- metrics는 slice(arm × split × subgroup 수준 × outcome × update)마다 한 번 정해서 그 "
        "slice의 모든 지표 행에 똑같이 적용한다. 회귀 cut 지표는 그 cut의 사건 수를 따로 본다.",
        "- 가린 건수에서 나온 비율, 지표, 곡선 점은 비운다(NaN). 임계값 격자의 한 점은 "
        "tp/fp/tn/fn 중 하나라도 가려지면 그 점 전체를 비운다.",
        f"- 사건이나 비사건이 {bundle.min_cell}건 미만인 slice의 지표와 곡선은 값을 비운다.",
        "- 연속형은 최솟값·최댓값 없이 평균, SD, 중앙값, 사분위, p5–p95만 낸다. 값이 "
        f"{bundle.min_cell}개 미만인 열은 비우고, 더해지는 열들(overall = dev + test, test = 결과군 "
        f"열 + 라벨 없는 ECG) 중 한 부분이 1–{bundle.min_cell - 1}개이면 그 변수의 연속형 요약을 "
        "모든 열에서 비운다(n × 평균의 차로 그 부분의 평균이 나오므로). Table 1의 `unit` 열은 그 행이 "
        "사람(person)을 세는지 ECG를 세는지 적는다.",
        "- 단계 하나의 제외 수 합계가 1–"
        f"{bundle.min_cell - 1}건이면 반출을 멈춘다(`--allow-small-stage-totals`로만 진행, "
        "export_manifest.json에 기록).",
        "",
        "## 가림으로 막지 않는 것",
        "",
        "- 임계값 격자(curves의 threshold_grid)는 0.01 간격의 누적 건수다. 이웃한 두 점의 차는 그 "
        "예측 위험 구간에 든 건수이고, 위험 분포(risk_distribution)의 구간 건수도 이렇게 격자에서 "
        "계산할 수 있다. 예측 위험 구간은 환자 속성이 아니므로 설계상 허용한다. 위험 분포의 가림은 "
        "표 사이의 일관성을 위한 것이지, 이 구간 건수를 숨기지는 않는다.",
        "- Kaplan-Meier 곡선의 계단 높이와 at_risk로 그 시점의 사건 수를 대략 알 수 있다(보통의 KM "
        "보고와 같다).",
        "- Table 1 건수와 metrics n이 둘 다 보이면 그 차는 그 split·수준에서 라벨이 없거나 예측이 "
        "없어 평가하지 않은 ECG 수의 상한이다. 라벨 없는 ECG 수는 위 규칙대로 보호한다(그 상한이 "
        f"1–{bundle.min_cell - 1}이면 한쪽을 가린다). 라벨은 있는데 그 arm이 예측하지 않은 ECG 수"
        "(라벨 있는 ECG − arm n)는 환자 속성이 아니라 따로 가리지 않는다.",
        "- retrieve·profile 요약에는 건수를 하나도 넣지 않는다(라벨 값 분위수만). 건수는 Table 1, "
        "exclusions, STARD에서 가린 대로만 나간다.",
        "",
        "## 사용한 기준",
        "",
        f"- min_cell: {bundle.min_cell}",
    ]
    if len(set(min_cells.values())) > 1:
        detail = ", ".join(f"{sid} {value}" for sid, value in sorted(min_cells.items()))
        lines.append(f"- study마다 `export.min_cell`이 달라({detail}) 가장 큰 값 {bundle.min_cell}을 썼다.")
    for sid, facts in sorted(by_study.items()):
        crit = facts["criteria"]
        grid = crit["net_benefit_grid"]
        lines.append(
            f"- {sid}: age_bands {', '.join('%g' % b for b in crit['age_bands'])}, "
            f"regression_cuts {cut_text(crit['regression_cuts']) or 'none'} "
            f"({crit['regression_cuts_from']}), "
            f"net_benefit_grid {grid['start']:g}–{grid['stop']:g} step {grid['step']:g} "
            f"({grid['n_points']}점), evaluate.min_events {crit['min_events']}"
            + (f", follow_up_after_last_visit_days {crit['follow_up_after_last_visit_days']} "
               f"({crit['follow_up_after_last_visit_days_from']})"
               if "follow_up_after_last_visit_days" in crit else "")
        )
    if not _has_openpyxl():
        lines += ["", "openpyxl이 없어 xlsx 파일(table1/*.xlsx, metrics/all_metrics.xlsx)을 만들지 않았다."]
    if not _has_matplotlib():
        lines += ["", "matplotlib이 없어 figures/를 만들지 않았다."]
    if bundle.figure_notes:
        lines += ["", "## 그림 메모", ""] + [f"- {note}" for note in bundle.figure_notes]
    return "\n".join(lines) + "\n"


def _source_sha256() -> Optional[str]:
    from mival.pipeline.manifest import code_identity

    return code_identity().get("source_sha256")


def build_bundle(studies: Sequence[str], study_root: Path, runs_root: Path, home: Path,
                 site: str, out: Path, allow_small_stage_totals: bool = False) -> Path:
    """Write the results zip for ``studies`` at ``out``; on any failure no zip is written.

    ``studies`` name directories under ``study_root`` holding ``study.yaml``;
    each study's ``runs.json`` is read from ``runs_root/<study_id>/_run/``.
    A stage that excluded 1 to ``min_cell - 1`` records in total fails the
    export unless ``allow_small_stage_totals``; the stages are then recorded
    in the export manifest and the README. A file already at ``out`` from an
    earlier export is left as it was when this one fails.
    """
    from mival.pipeline.study import load_study

    if not studies:
        raise ValueError("no studies to export")
    out = Path(out)
    loaded = [load_study(Path(study_root) / name) for name in studies]
    min_cells = {s.study_id: criteria.min_cell(s.stage_spec("export")) for s in loaded}
    min_cell = max(min_cells.values())
    for study in loaded:
        used = study_criteria(study)
        if used["min_events"] < min_cell:
            raise ValueError(
                f"study {study.study_id}: evaluate.min_events {used['min_events']} is below "
                f"export.min_cell {min_cell}; the evaluate stage would report curves for "
                "slices the bundle must mask. Set evaluate.min_events >= export.min_cell."
            )

    staging = Path(tempfile.mkdtemp(prefix="mival-export-"))
    partial = out.parent / f".{out.name}.partial"
    try:
        bundle = _Bundle(staging, min_cell, home, runs_root)
        bundle.allow_small_stage_totals = bool(allow_small_stage_totals)
        exported = {s.study_id: _export_study(bundle, s, Path(runs_root), Path(home)) for s in loaded}
        by_study = {sid: item["facts"] for sid, item in exported.items()}

        if _has_openpyxl():
            sheets = []
            for sid, item in exported.items():
                sheets += [(f"{sid}_metrics", item["metrics"]),
                           (f"{sid}_comparisons", item["comparisons"]),
                           (f"{sid}_sample_size", item["sample_size"])]
            _write_xlsx(bundle.path("metrics/all_metrics.xlsx"), sheets)

        _write_json(bundle.path(EXPORT_MANIFEST), {
            GENERATED_AT: datetime.now(timezone.utc).isoformat(timespec="seconds"),
            "site": site,
            "criteria": {"min_cell": min_cell,
                         "by_study": {sid: f["criteria"] for sid, f in by_study.items()}},
            "stages": {sid: item["stages"] for sid, item in exported.items()},
            # Per arm, the platform runtime override (e.g. "windows") its
            # models run used, so a pooled reader knows a site ran a card
            # differently from the reference machine.
            "arms": {sid: item["arms"] for sid, item in exported.items()},
            "export_source_sha256": _source_sha256(),
            "allow_small_stage_totals": bundle.allow_small_stage_totals,
            "small_stage_totals": bundle.small_stage_totals,
        })

        from mival.reporting import methods_draft, tripod_checklist

        bundle.path("reporting/methods_draft.md").write_text(
            methods_draft([by_study[s.study_id] for s in loaded]), encoding="utf-8")
        # The checklist and README list the bundle's files, themselves included.
        checklist = bundle.path("reporting/tripod_ai_checklist.md")
        readme = bundle.path("README.md")
        checklist.write_text("", encoding="utf-8")
        readme.write_text("", encoding="utf-8")
        checklist.write_text(tripod_checklist(bundle.files()), encoding="utf-8")
        readme.write_text(_readme(bundle, site, by_study, min_cells), encoding="utf-8")

        names = bundle.files()
        stray = [n for n in names if not any(re.fullmatch(p, n) for p in ALLOWLIST)]
        if stray:
            raise DisclosureError(f"not on the export allowlist: {', '.join(stray)}")
        check_no_patient_level([staging / n for n in names])

        out.parent.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(partial, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for name in names:
                archive.write(staging / name, name)
        os.replace(partial, out)
        return out
    finally:
        shutil.rmtree(staging, ignore_errors=True)
        if partial.exists():
            partial.unlink()


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def _today() -> str:
    return datetime.now(timezone.utc).strftime("%y%m%d")


def _studies_with_runs(study_root: Path, runs_root: Path) -> List[str]:
    from mival.pipeline.study import load_study

    names = []
    for path in sorted(Path(study_root).glob("*/study.yaml")):
        study = load_study(path)
        if (Path(runs_root) / study.study_id / RUNS_FILE).is_file():
            names.append(path.parent.name)
    return names


def main(argv: Optional[Sequence[str]] = None) -> int:
    from mival.pipeline.study import load_study

    parser = argparse.ArgumentParser(prog="python -m mival.export", description=__doc__.split("\n")[0])
    parser.add_argument("--studies", nargs="*", default=None,
                        help="study directories under --study-root (default: every one with runs.json)")
    parser.add_argument("--study-root", default="studies")
    parser.add_argument("--runs-root", default=os.environ.get("MIVAL_RUNS"))
    parser.add_argument("--home", default=os.environ.get("MIVAL_HOME"))
    parser.add_argument("--out", default=None, help="zip path (default results-<site>-<YYMMDD>.zip)")
    parser.add_argument("--allow-small-stage-totals", action="store_true",
                        help="export although a stage excluded 1..min_cell-1 records in total "
                             "(recorded in the bundle)")
    args = parser.parse_args(argv)
    if not args.home:
        parser.error("--home or MIVAL_HOME is required")
    runs_root = Path(args.runs_root) if args.runs_root else Path(args.home) / "runs"
    studies = list(args.studies) if args.studies else _studies_with_runs(Path(args.study_root), runs_root)
    if not studies:
        parser.error(f"no study under {args.study_root} has a runs.json under {runs_root}")
    sites = sorted({load_study(Path(args.study_root) / name).site for name in studies})
    if len(sites) != 1:
        parser.error(f"the studies name different sites: {', '.join(sites)}")
    out = Path(args.out) if args.out else Path(f"results-{sites[0]}-{_today()}.zip")
    try:
        path = build_bundle(studies, Path(args.study_root), runs_root, Path(args.home), sites[0],
                            out, allow_small_stage_totals=args.allow_small_stage_totals)
    except Exception:
        if out.exists():
            print(f"export failed; {out.resolve()} is from an earlier export, not this run",
                  file=sys.stderr)
        raise
    print(path.resolve())
    return 0


if __name__ == "__main__":  # pragma: no cover - exercised through main()
    raise SystemExit(main())
